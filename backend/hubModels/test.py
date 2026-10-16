import datetime
from unittest import mock

from _decimal import Decimal
from dateutil.relativedelta import relativedelta
from django.test import TestCase
from rest_framework.test import APIClient
from .factory import TransactionFactory

from .models import (HubUser, CustomLabel, Transaction, TransactionRecurrency, SavingPlan,
                     CategoryBudget, SavingPlanContribution, DebtLoan, DebtPayment)



class TransactionTestCase(TestCase):
    STARTER_WALLET_AMOUNT = 1000

    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        # Adding funds to user wallet
        self.wallet.update_balance(self.STARTER_WALLET_AMOUNT)
        
    def test_transaction_updates_wallet(self):
        """
        Check if a given transaction of type `EXPENSE`, when `update_wallet` is True, updates the wallet's current
        amount
        """
        transaction = self.factory.create_transaction_filled(self.wallet, type='EXPENSE', update_wallet=True)
        current_wallet_amount = self.STARTER_WALLET_AMOUNT - transaction.value
        self.assertEqual(Decimal(current_wallet_amount), self.wallet.get_current_amount())

    def test_transaction_is_from_wallet(self):
        transaction = self.factory.create_transaction_filled(self.wallet)
        self.assertEqual(transaction.is_from_wallet(self.wallet), True)

    def test_recurrent_transaction(self):
        """
        Check if creating a recurrent transaction successfully instantiates a `TransactionRecurrency` object with
        correct values using `create_from_json` function
        """
        transaction = self.factory.create_recurrent_transaction(self.user.pk, amount=2, duration="WEEKS")
        self.assertEqual(transaction.recurrent, True)

        recurrency = transaction.recurrency
        self.assertIsNotNone(recurrency)
        self.assertEqual(recurrency.amount, 2)
        self.assertEqual(recurrency.duration, "WEEKS")

    def test_get_monthly_earnings(self):
        current_month_start = datetime.date.today().replace(day=1)
        out_of_range = current_month_start - relativedelta(days=1)
        in_range = current_month_start + relativedelta(days=1)

        earnings_in_range = self.factory.create_transaction_filled(self.wallet, type="EARNING", value=150)
        earning_in_range2 = self.factory.create_transaction_filled(self.wallet, type="EARNING", value=50, date=in_range)
        earning_in_range3 = self.factory.create_transaction_filled(self.wallet, type="EARNING", value=300,
                                                                  date=current_month_start)

        earning_out_of_range = self.factory.create_transaction_filled(self.wallet, type="EARNING", date=out_of_range)

        # these transactions shouldn't be considered in the queryset since it's not an earning
        expense_in_range = self.factory.create_transaction_filled(self.wallet, type="EXPENSE", date=in_range)

        # Sum of all set transactions values
        expected_value = 500
        monthly_earning = self.wallet.get_monthly_earnings()
        self.assertEqual(expected_value, monthly_earning)

    def test_get_monthly_expenses(self):
        current_month_start = datetime.date.today().replace(day=1)
        out_of_range = current_month_start - relativedelta(days=1)
        in_range = current_month_start + relativedelta(days=1)

        expense_in_range = self.factory.create_transaction_filled(self.wallet, type="EXPENSE", value=150)
        expense_in_range2 = self.factory.create_transaction_filled(self.wallet, type="EXPENSE", value=50, date=in_range)
        expense_in_range3 = self.factory.create_transaction_filled(self.wallet, type="EXPENSE", value=300,
                                                                   date=current_month_start)
        expense_out_of_range = self.factory.create_transaction_filled(self.wallet, type="EXPENSE",date=out_of_range)

        # this transaction shouldn't be considered in the queryset since it's an earning
        earning_in_range = self.factory.create_transaction_filled(self.wallet, type="EARNING", date=in_range)

        expected_value = 500
        monthly_earning = self.wallet.get_monthly_expenses()
        self.assertEqual(expected_value, monthly_earning)


class TransactionAPIConsistencyTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.wallet.set_balance(1000)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_datetime_input_is_accepted_without_an_obsolete_description(self):
        transaction = Transaction.create_from_json({
            'title': 'Salary',
            'value': 125,
            'type': 'EARNING',
            'date': datetime.datetime.now(),
            'updateWallet': True,
        }, self.user.pk)

        self.wallet.refresh_from_db()
        self.assertEqual(transaction.date, datetime.date.today())
        self.assertEqual(self.wallet.current_amount, Decimal('1125.00'))

    def test_edit_reverses_previous_effect_before_applying_new_effect(self):
        transaction = self.factory.create_transaction_filled(
            self.wallet, title='Groceries', value=100, type='EXPENSE', update_wallet=True)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.current_amount, Decimal('900.00'))

        response = self.client.put(f'/api/transaction/{transaction.pk}', {
            'title': 'Refund',
            'value': 250,
            'type': 'EARNING',
            'date': '2026-10-01',
        }, format='json')

        self.assertEqual(response.status_code, 200)
        self.wallet.refresh_from_db()
        transaction.refresh_from_db()
        self.assertEqual(transaction.date, datetime.date(2026, 10, 1))
        self.assertEqual(self.wallet.current_amount, Decimal('1250.00'))

    def test_delete_reverses_effect_once_and_missing_or_foreign_records_are_not_server_errors(self):
        transaction = self.factory.create_transaction_filled(
            self.wallet, value=75, type='EXPENSE', update_wallet=True)

        response = self.client.delete(f'/api/transaction/{transaction.pk}')
        self.assertEqual(response.status_code, 200)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.current_amount, Decimal('1000.00'))

        response = self.client.delete(f'/api/transaction/{transaction.pk}')
        self.assertEqual(response.status_code, 404)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.current_amount, Decimal('1000.00'))

        other_user = self.factory.create_user()
        other_transaction = self.factory.create_transaction_filled(
            other_user.get_wallet(), value=20, type='EXPENSE', update_wallet=True)
        response = self.client.delete(f'/api/transaction/{other_transaction.pk}')
        self.assertEqual(response.status_code, 404)

    def test_invalid_transaction_input_returns_400(self):
        response = self.client.delete('/api/transaction/')
        self.assertEqual(response.status_code, 400)

        response = self.client.post('/api/transaction/', {
            'title': 'Invalid',
            'value': '-10',
            'type': 'EXPENSE',
            'date': 'not-a-date',
            'updateWallet': True,
        }, format='json')

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Transaction.objects.count(), 0)

        response = self.client.put('/api/transaction/999999', {
            'title': 'Missing',
        }, format='json')
        self.assertEqual(response.status_code, 404)

    def test_labels_are_validated_and_scoped_to_the_current_user(self):
        response = self.client.delete('/api/label/')
        self.assertEqual(response.status_code, 400)

        response = self.client.post('/api/label/', {
            'name': '',
            'color': 'not-a-colour',
        }, format='json')
        self.assertEqual(response.status_code, 400)

        other_user = self.factory.create_user()
        other_label = CustomLabel.objects.create(
            wallet=other_user.get_wallet(), name='Private', color='#000000')
        response = self.client.delete(f'/api/label/{other_label.pk}')
        self.assertEqual(response.status_code, 404)


class SavingPlanAPITestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_create_and_list_saving_plans(self):
        response = self.client.post('/api/savingplan/', {
            'title': 'New Laptop',
            'amount': '1500.00',
            'goal_date': '2026-12-31',
            'description': 'Saving for work machine',
        }, format='json')
        self.assertEqual(response.status_code, 201)
        plan_id = response.data.get('id')
        self.assertIsNotNone(plan_id)
        self.assertEqual(response.data['title'], 'New Laptop')
        self.assertEqual(response.data['active'], True)

        # List plans
        response = self.client.get('/api/savingplans/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]['id'], plan_id)

    def test_saving_plan_validation(self):
        # Missing title
        response = self.client.post('/api/savingplan/', {
            'title': '',
            'amount': 500,
        }, format='json')
        self.assertEqual(response.status_code, 400)

        # Invalid amount (negative)
        response = self.client.post('/api/savingplan/', {
            'title': 'Car',
            'amount': -100,
        }, format='json')
        self.assertEqual(response.status_code, 400)

        # Invalid date
        response = self.client.post('/api/savingplan/', {
            'title': 'Car',
            'amount': 100,
            'goal_date': 'invalid-date',
        }, format='json')
        self.assertEqual(response.status_code, 400)

    def test_get_detail_and_update_saving_plan(self):
        plan = self.user.get_wallet().saving_plans.create(
            title='Vacation',
            amount=Decimal('2000.00'),
            goal_date=datetime.date(2026, 11, 1),
            description='Beach trip'
        )

        # Get details
        response = self.client.get(f'/api/savingplan/{plan.pk}')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['title'], 'Vacation')

        # Update details
        response = self.client.put(f'/api/savingplan/{plan.pk}', {
            'title': 'Europe Vacation',
            'amount': '3000.00',
        }, format='json')
        self.assertEqual(response.status_code, 200)
        plan.refresh_from_db()
        self.assertEqual(plan.title, 'Europe Vacation')
        self.assertEqual(plan.amount, Decimal('3000.00'))

    def test_toggle_active_and_delete_saving_plan(self):
        plan = self.user.get_wallet().saving_plans.create(
            title='Emergency Fund',
            amount=Decimal('5000.00'),
            active=True
        )

        # Toggle / Deactivate plan via patch
        response = self.client.patch(f'/api/savingplan/{plan.pk}', {
            'active': False
        }, format='json')
        self.assertEqual(response.status_code, 200)
        plan.refresh_from_db()
        self.assertFalse(plan.active)

        # Toggle back
        response = self.client.patch(f'/api/savingplan/{plan.pk}')
        self.assertEqual(response.status_code, 200)
        plan.refresh_from_db()
        self.assertTrue(plan.active)

        # Delete plan
        response = self.client.delete(f'/api/savingplan/{plan.pk}')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.user.get_wallet().saving_plans.count(), 0)

    def test_saving_plan_isolation_between_users(self):
        other_user = self.factory.create_user()
        other_plan = other_user.get_wallet().saving_plans.create(
            title='Secret Fund',
            amount=Decimal('9999.00')
        )

        # Trying to view other user's plan
        response = self.client.get(f'/api/savingplan/{other_plan.pk}')
        self.assertEqual(response.status_code, 404)

        # Trying to edit other user's plan
        response = self.client.put(f'/api/savingplan/{other_plan.pk}', {
            'title': 'Hacked',
        }, format='json')
        self.assertEqual(response.status_code, 404)

        # Trying to delete other user's plan
        response = self.client.delete(f'/api/savingplan/{other_plan.pk}')
        self.assertEqual(response.status_code, 404)


class SavingPlanAdvancedTrackingTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

        self.plan = self.wallet.saving_plans.create(
            title='House Deposit',
            amount=Decimal('10000.00'),
            goal_date=datetime.date(2027, 12, 31),
            description='Saving for first apartment'
        )

    def test_goal_progress_serialization(self):
        # Initial state: 0 contributions
        response = self.client.get(f'/api/savingplan/{self.plan.pk}')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Decimal(str(response.data['amount_saved'])), Decimal('0.00'))
        self.assertEqual(Decimal(str(response.data['remaining_amount'])), Decimal('10000.00'))
        self.assertEqual(response.data['completion_percentage'], 0.0)
        self.assertEqual(response.data['status'], 'active')

        # Add contribution
        self.client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '2500.00',
            'date': '2026-10-01',
            'description': 'Initial bonus'
        }, format='json')

        response = self.client.get(f'/api/savingplan/{self.plan.pk}')
        self.assertEqual(Decimal(str(response.data['amount_saved'])), Decimal('2500.00'))
        self.assertEqual(Decimal(str(response.data['remaining_amount'])), Decimal('7500.00'))
        self.assertEqual(response.data['completion_percentage'], 25.0)
        self.assertEqual(response.data['status'], 'active')

        # Add more contribution to reach 100%
        self.client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '7500.00',
            'date': '2026-10-10',
            'description': 'Final payment'
        }, format='json')

        response = self.client.get(f'/api/savingplan/{self.plan.pk}')
        self.assertEqual(Decimal(str(response.data['amount_saved'])), Decimal('10000.00'))
        self.assertEqual(Decimal(str(response.data['remaining_amount'])), Decimal('0.00'))
        self.assertEqual(response.data['completion_percentage'], 100.0)
        self.assertEqual(response.data['status'], 'completed')

    def test_contribution_crud_and_idempotency(self):
        # 1. Add contribution with idempotency key
        key = 'unique-contrib-key-123'
        resp1 = self.client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '1000.00',
            'date': '2026-10-05',
            'description': 'Monthly saving',
            'idempotency_key': key,
        }, format='json')
        self.assertEqual(resp1.status_code, 201)
        contrib_id = resp1.data['id']

        # 2. Resubmit with identical idempotency key -> duplicate prevented
        resp2 = self.client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '1000.00',
            'date': '2026-10-05',
            'description': 'Monthly saving',
            'idempotency_key': key,
        }, format='json')
        self.assertEqual(resp2.status_code, 200)
        self.assertEqual(resp2.data['id'], contrib_id)
        self.assertEqual(self.plan.contributions.count(), 1)

        # 3. List contributions
        list_resp = self.client.get(f'/api/savingplans/{self.plan.pk}/contributions/')
        self.assertEqual(list_resp.status_code, 200)
        self.assertEqual(len(list_resp.data), 1)

        # 4. Update contribution
        update_resp = self.client.put(f'/api/savingplans/{self.plan.pk}/contributions/{contrib_id}', {
            'amount': '1200.00',
            'description': 'Adjusted monthly saving',
        }, format='json')
        self.assertEqual(update_resp.status_code, 200)
        self.assertEqual(Decimal(str(update_resp.data['amount'])), Decimal('1200.00'))

        # 5. Delete contribution
        del_resp = self.client.delete(f'/api/savingplans/{self.plan.pk}/contributions/{contrib_id}')
        self.assertEqual(del_resp.status_code, 200)
        self.assertEqual(self.plan.contributions.count(), 0)

    def test_contribution_validation(self):
        # Negative amount
        resp = self.client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '-500.00',
        }, format='json')
        self.assertEqual(resp.status_code, 400)

        # Zero amount
        resp = self.client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '0.00',
        }, format='json')
        self.assertEqual(resp.status_code, 400)

        # Invalid date
        resp = self.client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '500.00',
            'date': 'invalid-date',
        }, format='json')
        self.assertEqual(resp.status_code, 400)

    def test_contribution_cross_user_isolation(self):
        other_user = self.factory.create_user()
        other_client = APIClient()
        other_client.force_authenticate(other_user)

        # Other user cannot add contribution to self.plan
        resp = other_client.post(f'/api/savingplans/{self.plan.pk}/contributions/', {
            'amount': '500.00',
        }, format='json')
        self.assertEqual(resp.status_code, 404)

        # Create a contribution by owner
        contrib = self.plan.contributions.create(
            user=self.user,
            amount=Decimal('500.00'),
            date=datetime.date.today()
        )

        # Other user cannot view contributions
        resp = other_client.get(f'/api/savingplans/{self.plan.pk}/contributions/')
        self.assertEqual(resp.status_code, 404)

        # Other user cannot edit contribution
        resp = other_client.put(f'/api/savingplans/{self.plan.pk}/contributions/{contrib.pk}', {
            'amount': '1000.00',
        }, format='json')
        self.assertEqual(resp.status_code, 404)

        # Other user cannot delete contribution
        resp = other_client.delete(f'/api/savingplans/{self.plan.pk}/contributions/{contrib.pk}')
        self.assertEqual(resp.status_code, 404)

    def test_completion_estimate(self):
        # 1. No contributions -> Insufficient history
        est1 = self.plan.get_estimated_completion()
        self.assertIsNone(est1['estimated_date'])
        self.assertIn('Chưa có lịch sử', est1['message'])

        # 2. Contributions spanning multiple days
        today = datetime.date.today()
        # Contribution 1: 10 days ago, 1000
        self.plan.contributions.create(
            user=self.user,
            amount=Decimal('1000.00'),
            date=today - datetime.timedelta(days=10)
        )
        # Contribution 2: 5 days ago, 1000
        self.plan.contributions.create(
            user=self.user,
            amount=Decimal('1000.00'),
            date=today - datetime.timedelta(days=5)
        )
        # Total contributed: 2000 in 10 days => daily_rate = 200/day
        # Remaining: 8000 => 8000 / 200 = 40 days needed
        est2 = self.plan.get_estimated_completion()
        self.assertIsNotNone(est2['estimated_date'])
        self.assertEqual(est2['days_remaining'], 40)
        self.assertEqual(est2['daily_rate'], 200.0)

        # 3. Already completed plan
        self.plan.contributions.create(
            user=self.user,
            amount=Decimal('8000.00'),
            date=today
        )
        est3 = self.plan.get_estimated_completion()
        self.assertEqual(est3['status'], 'completed')
        self.assertIn('hoàn thành', est3['message'].lower())


class RecurringTransactionsTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.wallet.set_balance(1000)

    def test_recurring_generation_and_idempotency(self):
        from .recurring import process_due_recurring_transactions

        # Create a monthly recurring transaction starting 2 months ago
        start_date = datetime.date(2026, 8, 1)
        base_tx = Transaction.objects.create(
            wallet=self.wallet,
            title='Gym Subscription',
            value=Decimal('50.00'),
            type='EXPENSE',
            date=start_date,
            update_wallet=True,
            recurrent=True,
        )
        TransactionRecurrency.objects.create(
            transaction=base_tx,
            amount=1,
            duration='MONTHS',
            active=True
        )
        base_tx.recurrency = TransactionRecurrency.objects.get(transaction=base_tx)
        base_tx.save()

        # Target date is 2026-10-12
        as_of = datetime.date(2026, 10, 12)
        result = process_due_recurring_transactions(as_of_date=as_of)

        # Due dates: 2026-09-01 and 2026-10-01 (2 occurrences)
        self.assertEqual(result['generated_count'], 2)
        generated = Transaction.objects.filter(base_transaction=base_tx)
        self.assertEqual(generated.count(), 2)

        # Wallet should be updated: 1000 - 50 - 50 = 900
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.current_amount, Decimal('900.00'))

        # Run again with same as_of date: MUST be idempotent (0 new transactions generated)
        result2 = process_due_recurring_transactions(as_of_date=as_of)
        self.assertEqual(result2['generated_count'], 0)
        self.assertEqual(Transaction.objects.filter(base_transaction=base_tx).count(), 2)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.current_amount, Decimal('900.00'))

    def test_recurring_with_end_date(self):
        from .recurring import process_due_recurring_transactions

        from django.utils import timezone
        start_date = datetime.date(2026, 9, 1)
        end_date = timezone.make_aware(datetime.datetime(2026, 9, 20, 23, 59, 59))
        base_tx = Transaction.objects.create(
            wallet=self.wallet,
            title='Weekly Course',
            value=Decimal('20.00'),
            type='EXPENSE',
            date=start_date,
            update_wallet=False,
            recurrent=True,
        )
        rec = TransactionRecurrency.objects.create(
            transaction=base_tx,
            amount=1,
            duration='WEEKS',
            end_date=end_date,
            active=True
        )

        as_of = datetime.date(2026, 10, 12)
        result = process_due_recurring_transactions(as_of_date=as_of)

        # Week 1: Sep 8 (generated), Week 2: Sep 15 (generated), Week 3: Sep 22 (> Sep 20 end_date -> stops & deactivates)
        self.assertEqual(result['generated_count'], 2)
        rec.refresh_from_db()
        self.assertFalse(rec.active)


class TransactionSearchAndSerializationTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_search_transactions_by_title_label_type_amount(self):
        food_label = CustomLabel.objects.create(wallet=self.wallet, name='Dining', color='#ff0000')
        t1 = self.factory.create_transaction_filled(
            self.wallet, title='Coffee at Starbucks', value=55, type='EXPENSE', label=food_label, date=datetime.date(2026, 10, 1))
        t2 = self.factory.create_transaction_filled(
            self.wallet, title='Freelance Payment', value=800, type='EARNING', date=datetime.date(2026, 10, 2))
        t3 = self.factory.create_transaction_filled(
            self.wallet, title='Electric Bill', value=120, type='EXPENSE', date=datetime.date(2026, 10, 3))

        # Search by title keyword
        res = self.client.get('/api/transactions/?search=Coffee')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data), 1)
        self.assertEqual(res.data[0]['title'], 'Coffee at Starbucks')

        # Search by label name
        res = self.client.get('/api/transactions/?search=Dining')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data), 1)
        self.assertEqual(res.data[0]['id'], t1.pk)

        # Search by type keyword (EARNING / Thu)
        res = self.client.get('/api/transactions/?search=EARNING')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data), 1)
        self.assertEqual(res.data[0]['id'], t2.pk)

        # Search by amount
        res = self.client.get('/api/transactions/?search=120')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data), 1)
        self.assertEqual(res.data[0]['id'], t3.pk)

    def test_recurring_transaction_serialization(self):
        from hubModels.serializers import TransactionSerializer

        t_recurrent = self.factory.create_recurrent_transaction(self.user.pk, amount=3, duration='MONTHS')
        serializer = TransactionSerializer(t_recurrent)
        self.assertTrue(serializer.data['recurrent'])
        self.assertEqual(serializer.data['amount'], 3)
        self.assertEqual(serializer.data['duration'], 'MONTHS')

        # Non-recurrent transaction
        t_regular = self.factory.create_transaction_filled(self.wallet)
        serializer_reg = TransactionSerializer(t_regular)
        self.assertFalse(serializer_reg.data['recurrent'])
        self.assertIsNone(serializer_reg.data['amount'])
        self.assertIsNone(serializer_reg.data['duration'])

    def test_get_transactions_with_limit_and_chart_type_zero(self):
        for i in range(5):
            self.factory.create_transaction_filled(
                self.wallet, title=f'Trans {i}', value=10 * (i + 1), type='EXPENSE', date=datetime.date(2026, 10, i + 1)
            )

        res = self.client.get('/api/transactions/?limit=3&chart_type=0')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data), 3)
        self.assertEqual(res.data[0]['title'], 'Trans 4')
        self.assertEqual(res.data[1]['title'], 'Trans 3')
        self.assertEqual(res.data[2]['title'], 'Trans 2')


class CategoryBudgetModelTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.food_label = CustomLabel.objects.create(wallet=self.wallet, name='Food', color='#2ecc71')
        self.bills_label = CustomLabel.objects.create(wallet=self.wallet, name='Bills', color='#e74c3c')

    def test_create_budget_and_uniqueness_constraint(self):
        budget = CategoryBudget.create_from_json({
            'label': self.food_label.pk,
            'amount': 2000000,
            'month': 10,
            'year': 2026,
        }, user_pk=self.user.pk)

        self.assertEqual(budget.amount, Decimal('2000000.00'))
        self.assertEqual(budget.month, 10)
        self.assertEqual(budget.year, 2026)
        self.assertIsNotNone(budget.created_at)
        self.assertIsNotNone(budget.updated_at)

        # Duplicate budget for same wallet, label, month, year must fail
        from django.core.exceptions import ValidationError
        with self.assertRaises(ValidationError):
            CategoryBudget.create_from_json({
                'label': self.food_label.pk,
                'amount': 3000000,
                'month': 10,
                'year': 2026,
            }, user_pk=self.user.pk)

    def test_budget_validations(self):
        from django.core.exceptions import ValidationError

        # Amount must be positive
        with self.assertRaises(ValidationError):
            CategoryBudget.create_from_json({
                'label': self.food_label.pk,
                'amount': -100,
                'month': 10,
                'year': 2026,
            }, user_pk=self.user.pk)

        # Month must be between 1 and 12
        with self.assertRaises(ValidationError):
            CategoryBudget.create_from_json({
                'label': self.food_label.pk,
                'amount': 500000,
                'month': 13,
                'year': 2026,
            }, user_pk=self.user.pk)

        # Label from another user's wallet
        other_user = self.factory.create_user()
        other_label = CustomLabel.objects.create(wallet=other_user.get_wallet(), name='Other', color='#000000')
        with self.assertRaises(ValidationError):
            CategoryBudget.create_from_json({
                'label': other_label.pk,
                'amount': 500000,
                'month': 10,
                'year': 2026,
            }, user_pk=self.user.pk)

    def test_budget_spent_calculation_and_over_budget(self):
        budget = CategoryBudget.objects.create(
            wallet=self.wallet,
            label=self.food_label,
            amount=Decimal('1000000.00'),
            month=10,
            year=2026,
        )

        # In-scope expenses
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=300000, date=datetime.date(2026, 10, 5)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=400000, date=datetime.date(2026, 10, 10)
        )

        # Out-of-scope: different month
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=200000, date=datetime.date(2026, 9, 15)
        )

        # Out-of-scope: different label
        self.factory.create_transaction_filled(
            self.wallet, label=self.bills_label, type='EXPENSE', value=500000, date=datetime.date(2026, 10, 12)
        )

        # Out-of-scope: earning transaction with same label
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EARNING', value=100000, date=datetime.date(2026, 10, 8)
        )

        # Total spent should only be 300,000 + 400,000 = 700,000
        self.assertEqual(budget.get_spent(), Decimal('700000.00'))
        self.assertEqual(budget.get_remaining(), Decimal('300000.00'))
        self.assertEqual(budget.get_percentage(), 70.0)
        self.assertFalse(budget.is_over_budget())

        # Exceed budget
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=400000, date=datetime.date(2026, 10, 20)
        )
        self.assertEqual(budget.get_spent(), Decimal('1100000.00'))
        self.assertEqual(budget.get_remaining(), Decimal('0.00'))
        self.assertEqual(budget.get_percentage(), 110.0)
        self.assertTrue(budget.is_over_budget())


class CategoryBudgetAPITestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

        self.food_label = CustomLabel.objects.create(wallet=self.wallet, name='Food', color='#2ecc71')
        self.util_label = CustomLabel.objects.create(wallet=self.wallet, name='Utilities', color='#3498db')
        self.shop_label = CustomLabel.objects.create(wallet=self.wallet, name='Shopping', color='#e74c3c')

    def test_list_budgets_for_month_with_calculated_fields_and_statuses(self):
        # 1. Normal status (50%)
        b_food = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.food_label, amount=Decimal('1000000.00'), month=10, year=2026
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=500000, date=datetime.date(2026, 10, 2)
        )

        # 2. Approaching status (85%)
        b_util = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.util_label, amount=Decimal('1000000.00'), month=10, year=2026
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.util_label, type='EXPENSE', value=850000, date=datetime.date(2026, 10, 5)
        )

        # 3. Exceeded status (120%)
        b_shop = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.shop_label, amount=Decimal('1000000.00'), month=10, year=2026
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.shop_label, type='EXPENSE', value=1200000, date=datetime.date(2026, 10, 10)
        )

        res = self.client.get('/api/budgets/?month=10&year=2026')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data), 3)

        budgets_by_label = {item['label']['name']: item for item in res.data}

        # Food assertions
        food_res = budgets_by_label['Food']
        self.assertEqual(food_res['limit'], '1000000.00')
        self.assertEqual(food_res['spent'], '500000.00')
        self.assertEqual(food_res['remaining'], '500000.00')
        self.assertEqual(food_res['percentage'], 50.0)
        self.assertEqual(food_res['status'], 'normal')
        self.assertFalse(food_res['is_over_budget'])

        # Utilities assertions
        util_res = budgets_by_label['Utilities']
        self.assertEqual(util_res['limit'], '1000000.00')
        self.assertEqual(util_res['spent'], '850000.00')
        self.assertEqual(util_res['remaining'], '150000.00')
        self.assertEqual(util_res['percentage'], 85.0)
        self.assertEqual(util_res['status'], 'approaching')
        self.assertFalse(util_res['is_over_budget'])

        # Shopping assertions
        shop_res = budgets_by_label['Shopping']
        self.assertEqual(shop_res['limit'], '1000000.00')
        self.assertEqual(shop_res['spent'], '1200000.00')
        self.assertEqual(shop_res['remaining'], '0.00')
        self.assertEqual(shop_res['percentage'], 120.0)
        self.assertEqual(shop_res['status'], 'exceeded')
        self.assertTrue(shop_res['is_over_budget'])

    def test_create_and_retrieve_budget_detail(self):
        # Create budget via POST
        create_res = self.client.post('/api/budgets/', {
            'label': self.food_label.pk,
            'amount': 2500000,
            'month': 11,
            'year': 2026,
        }, format='json')
        self.assertEqual(create_res.status_code, 201)
        budget_id = create_res.data['id']
        self.assertEqual(create_res.data['limit'], '2500000.00')
        self.assertEqual(create_res.data['month'], 11)

        # Retrieve single detail via GET /api/budgets/<id>
        detail_res = self.client.get(f'/api/budgets/{budget_id}')
        self.assertEqual(detail_res.status_code, 200)
        self.assertEqual(detail_res.data['id'], budget_id)
        self.assertEqual(detail_res.data['label']['id'], self.food_label.pk)

        # Retrieve non-existent budget
        not_found_res = self.client.get('/api/budgets/99999')
        self.assertEqual(not_found_res.status_code, 404)

    def test_create_budget_validations_and_boundary_behavior(self):
        # Missing category
        res1 = self.client.post('/api/budgets/', {
            'amount': 500000,
            'month': 10,
            'year': 2026,
        }, format='json')
        self.assertEqual(res1.status_code, 400)

        # Zero amount
        res2 = self.client.post('/api/budgets/', {
            'label': self.food_label.pk,
            'amount': 0,
            'month': 10,
            'year': 2026,
        }, format='json')
        self.assertEqual(res2.status_code, 400)

        # Negative amount
        res3 = self.client.post('/api/budgets/', {
            'label': self.food_label.pk,
            'amount': -100000,
            'month': 10,
            'year': 2026,
        }, format='json')
        self.assertEqual(res3.status_code, 400)

        # Invalid month
        res4 = self.client.post('/api/budgets/', {
            'label': self.food_label.pk,
            'amount': 100000,
            'month': 0,
            'year': 2026,
        }, format='json')
        self.assertEqual(res4.status_code, 400)

        # Duplicate budget creation
        self.client.post('/api/budgets/', {
            'label': self.food_label.pk,
            'amount': 500000,
            'month': 10,
            'year': 2026,
        }, format='json')
        dup_res = self.client.post('/api/budgets/', {
            'label': self.food_label.pk,
            'amount': 700000,
            'month': 10,
            'year': 2026,
        }, format='json')
        self.assertEqual(dup_res.status_code, 400)

    def test_update_and_delete_budget(self):
        budget = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.food_label, amount=Decimal('500000.00'), month=10, year=2026
        )

        # Update budget
        update_res = self.client.put(f'/api/budgets/{budget.pk}', {
            'amount': 800000,
        }, format='json')
        self.assertEqual(update_res.status_code, 200)
        self.assertEqual(update_res.data['limit'], '800000.00')

        budget.refresh_from_db()
        self.assertEqual(budget.amount, Decimal('800000.00'))

        # Delete budget
        del_res = self.client.delete(f'/api/budgets/{budget.pk}')
        self.assertEqual(del_res.status_code, 200)
        self.assertFalse(CategoryBudget.objects.filter(pk=budget.pk).exists())

    def test_budget_summary_endpoint_and_user_isolation(self):
        # Set up budgets for current user
        CategoryBudget.objects.create(
            wallet=self.wallet, label=self.food_label, amount=Decimal('1000000.00'), month=10, year=2026
        )
        CategoryBudget.objects.create(
            wallet=self.wallet, label=self.util_label, amount=Decimal('2000000.00'), month=10, year=2026
        )

        # Expenses for current user
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=600000, date=datetime.date(2026, 10, 5)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.util_label, type='EXPENSE', value=1700000, date=datetime.date(2026, 10, 8)
        )
        # Earning with food label should not be subtracted from expense spending
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EARNING', value=500000, date=datetime.date(2026, 10, 10)
        )
        # Unbudgeted expense (shopping label has no budget for this month)
        self.factory.create_transaction_filled(
            self.wallet, label=self.shop_label, type='EXPENSE', value=200000, date=datetime.date(2026, 10, 12)
        )

        # Other user's budget and transactions
        other_user = self.factory.create_user()
        other_wallet = other_user.get_wallet()
        other_label = CustomLabel.objects.create(wallet=other_wallet, name='OtherFood', color='#111111')
        CategoryBudget.objects.create(
            wallet=other_wallet, label=other_label, amount=Decimal('5000000.00'), month=10, year=2026
        )
        self.factory.create_transaction_filled(
            other_wallet, label=other_label, type='EXPENSE', value=4000000, date=datetime.date(2026, 10, 5)
        )

        summary_res = self.client.get('/api/budgets/summary/?month=10&year=2026')
        self.assertEqual(summary_res.status_code, 200)
        data = summary_res.data
        self.assertEqual(data['month'], 10)
        self.assertEqual(data['year'], 2026)
        # Total budget: 1,000,000 + 2,000,000 = 3,000,000 (other user's 5,000,000 strictly excluded)
        self.assertEqual(data['total_budget'], '3000000.00')
        # Total spent in budgeted categories: 600,000 + 1,700,000 = 2,300,000
        self.assertEqual(data['total_spent'], '2300000.00')
        # Total remaining: 700,000
        self.assertEqual(data['total_remaining'], '700000.00')
        # Overall percentage: 2,300,000 / 3,000,000 = 76.67%
        self.assertEqual(data['overall_percentage'], 76.67)
        # Total monthly expenses across all categories: 2,300,000 + 200,000 = 2,500,000
        self.assertEqual(data['total_monthly_expenses'], '2500000.00')
        # Unbudgeted expenses: 200,000
        self.assertEqual(data['unbudgeted_spent'], '200000.00')
        self.assertEqual(data['budget_count'], 2)
        # Statuses: Food (60%) -> normal, Util (85%) -> approaching
        self.assertEqual(data['status_counts']['normal'], 1)
        self.assertEqual(data['status_counts']['approaching'], 1)
        self.assertEqual(data['status_counts']['exceeded'], 0)

    def test_user_ownership_and_unauthorized_access(self):
        other_user = self.factory.create_user()
        other_wallet = other_user.get_wallet()
        other_label = CustomLabel.objects.create(wallet=other_wallet, name='OtherLabel', color='#555555')
        other_budget = CategoryBudget.objects.create(
            wallet=other_wallet, label=other_label, amount=Decimal('500000.00'), month=10, year=2026
        )

        # 1. Current user cannot retrieve other user's budget detail
        res = self.client.get(f'/api/budgets/{other_budget.pk}')
        self.assertEqual(res.status_code, 404)

        # 2. Current user cannot update other user's budget
        update_res = self.client.put(f'/api/budgets/{other_budget.pk}', {'amount': 900000}, format='json')
        self.assertEqual(update_res.status_code, 404)

        # 3. Current user cannot delete other user's budget
        del_res = self.client.delete(f'/api/budgets/{other_budget.pk}')
        self.assertEqual(del_res.status_code, 404)
        self.assertTrue(CategoryBudget.objects.filter(pk=other_budget.pk).exists())

        # 4. Unauthenticated user gets 401
        anon_client = APIClient()
        anon_res = anon_client.get('/api/budgets/')
        self.assertEqual(anon_res.status_code, 401)

    def test_monthly_period_boundaries(self):
        # Budget for October 2026
        budget = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.food_label, amount=Decimal('1000000.00'), month=10, year=2026
        )

        # Transaction on last day of September (2026-09-30) -> MUST NOT be included
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=100000, date=datetime.date(2026, 9, 30)
        )
        # Transaction on first day of October (2026-10-01) -> MUST be included
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=200000, date=datetime.date(2026, 10, 1)
        )
        # Transaction on last day of October (2026-10-31) -> MUST be included
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=300000, date=datetime.date(2026, 10, 31)
        )
        # Transaction on first day of November (2026-11-01) -> MUST NOT be included
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=400000, date=datetime.date(2026, 11, 1)
        )

        res = self.client.get('/api/budgets/?month=10&year=2026')
        self.assertEqual(res.status_code, 200)
        food_budget_data = res.data[0]
        # Only 200,000 + 300,000 = 500,000 in October
        self.assertEqual(food_budget_data['spent'], '500000.00')
        self.assertEqual(food_budget_data['remaining'], '500000.00')
        self.assertEqual(food_budget_data['percentage'], 50.0)

    def test_expense_aggregation_and_different_categories(self):
        # Two category budgets
        b_food = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.food_label, amount=Decimal('1000000.00'), month=10, year=2026
        )
        b_util = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.util_label, amount=Decimal('800000.00'), month=10, year=2026
        )

        # Multiple transactions for Food on different dates
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=150000, date=datetime.date(2026, 10, 3)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=250000, date=datetime.date(2026, 10, 14)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=100000, date=datetime.date(2026, 10, 25)
        )

        # Multiple transactions for Utilities on different dates
        self.factory.create_transaction_filled(
            self.wallet, label=self.util_label, type='EXPENSE', value=300000, date=datetime.date(2026, 10, 5)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.util_label, type='EXPENSE', value=340000, date=datetime.date(2026, 10, 20)
        )

        # Income transactions (should NOT reduce expense totals)
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EARNING', value=500000, date=datetime.date(2026, 10, 15)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.util_label, type='EARNING', value=200000, date=datetime.date(2026, 10, 22)
        )

        # Unlabeled expense (should not affect either category)
        self.factory.create_transaction_filled(
            self.wallet, label=None, type='EXPENSE', value=99000, date=datetime.date(2026, 10, 18)
        )

        res = self.client.get('/api/budgets/?month=10&year=2026')
        self.assertEqual(res.status_code, 200)
        budgets_by_label = {item['label']['name']: item for item in res.data}

        # Food: 150,000 + 250,000 + 100,000 = 500,000 (50%)
        self.assertEqual(budgets_by_label['Food']['spent'], '500000.00')
        self.assertEqual(budgets_by_label['Food']['remaining'], '500000.00')
        self.assertEqual(budgets_by_label['Food']['status'], 'normal')

        # Utilities: 300,000 + 340,000 = 640,000 (640,000 / 800,000 = 80%) -> approaching
        self.assertEqual(budgets_by_label['Utilities']['spent'], '640000.00')
        self.assertEqual(budgets_by_label['Utilities']['remaining'], '160000.00')
        self.assertEqual(budgets_by_label['Utilities']['percentage'], 80.0)
        self.assertEqual(budgets_by_label['Utilities']['status'], 'approaching')

    def test_percentage_and_remaining_balance_edge_cases(self):
        # 1. Zero spending
        b_food = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.food_label, amount=Decimal('500000.00'), month=10, year=2026
        )
        res = self.client.get(f'/api/budgets/{b_food.pk}')
        self.assertEqual(res.data['spent'], '0.00')
        self.assertEqual(res.data['remaining'], '500000.00')
        self.assertEqual(res.data['percentage'], 0.0)
        self.assertEqual(res.data['status'], 'normal')
        self.assertFalse(res.data['is_over_budget'])

        # 2. Exact limit (100%)
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=500000, date=datetime.date(2026, 10, 5)
        )
        res_100 = self.client.get(f'/api/budgets/{b_food.pk}')
        self.assertEqual(res_100.data['spent'], '500000.00')
        self.assertEqual(res_100.data['remaining'], '0.00')
        self.assertEqual(res_100.data['percentage'], 100.0)
        self.assertEqual(res_100.data['status'], 'approaching')
        self.assertFalse(res_100.data['is_over_budget'])

        # 3. Exceeded (140%)
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=200000, date=datetime.date(2026, 10, 15)
        )
        res_140 = self.client.get(f'/api/budgets/{b_food.pk}')
        self.assertEqual(res_140.data['spent'], '700000.00')
        self.assertEqual(res_140.data['remaining'], '0.00')
        self.assertEqual(res_140.data['percentage'], 140.0)
        self.assertEqual(res_140.data['status'], 'exceeded')
        self.assertTrue(res_140.data['is_over_budget'])


class FinancialReportAPITestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

        self.food_label = CustomLabel.objects.create(wallet=self.wallet, name='Food', color='#2ecc71')
        self.bills_label = CustomLabel.objects.create(wallet=self.wallet, name='Bills', color='#e74c3c')
        self.salary_label = CustomLabel.objects.create(wallet=self.wallet, name='Salary', color='#3498db')

    def test_report_summary_calculations_and_filtering(self):
        # Current period: 2026-10-01 to 2026-10-10 (10 days)
        # Earnings
        self.factory.create_transaction_filled(
            self.wallet, label=self.salary_label, type='EARNING', value=10000000, date=datetime.date(2026, 10, 1)
        )
        # Expenses
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=1500000, date=datetime.date(2026, 10, 2)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.bills_label, type='EXPENSE', value=2500000, date=datetime.date(2026, 10, 5)
        )
        # Uncategorized expense
        self.factory.create_transaction_filled(
            self.wallet, label=None, type='EXPENSE', value=1000000, date=datetime.date(2026, 10, 8)
        )
        # Out-of-range transaction (2026-10-15)
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=500000, date=datetime.date(2026, 10, 15)
        )

        res = self.client.get('/api/reports/summary/?start_date=2026-10-01&end_date=2026-10-10')
        self.assertEqual(res.status_code, 200)
        data = res.data

        # Period check
        self.assertEqual(data['period']['days_count'], 10)
        self.assertEqual(data['period']['start_date'], '2026-10-01')
        self.assertEqual(data['period']['end_date'], '2026-10-10')

        # Summary check
        # Total income: 10,000,000
        self.assertEqual(data['summary']['total_income'], '10000000.00')
        # Total expense: 1,500,000 + 2,500,000 + 1,000,000 = 5,000,000
        self.assertEqual(data['summary']['total_expense'], '5000000.00')
        # Net cash flow: 10,000,000 - 5,000,000 = 5,000,000
        self.assertEqual(data['summary']['net_cash_flow'], '5000000.00')
        # Average daily spending: 5,000,000 / 10 = 500,000.00
        self.assertEqual(data['summary']['average_daily_spending'], '500000.00')

        # Category breakdown
        breakdown = {item['label_name']: item for item in data['category_breakdown']}
        self.assertIn('Food', breakdown)
        self.assertIn('Bills', breakdown)
        self.assertIn('Chưa phân loại', breakdown)
        self.assertEqual(breakdown['Bills']['total'], '2500000.00')
        self.assertEqual(breakdown['Bills']['percentage'], 50.0)
        self.assertEqual(breakdown['Food']['total'], '1500000.00')
        self.assertEqual(breakdown['Food']['percentage'], 30.0)
        self.assertEqual(breakdown['Chưa phân loại']['total'], '1000000.00')
        self.assertEqual(breakdown['Chưa phân loại']['percentage'], 20.0)

    def test_report_comparison_with_previous_period(self):
        # Previous period: 2026-10-01 to 2026-10-10 (10 days)
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=1000000, date=datetime.date(2026, 10, 5)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.salary_label, type='EARNING', value=6000000, date=datetime.date(2026, 10, 1)
        )

        # Current period: 2026-10-11 to 2026-10-20 (10 days)
        # Food increased from 1,000,000 to 2,500,000 (+1,500,000, +150%)
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=2500000, date=datetime.date(2026, 10, 15)
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.salary_label, type='EARNING', value=8000000, date=datetime.date(2026, 10, 12)
        )

        res = self.client.get('/api/reports/summary/?start_date=2026-10-11&end_date=2026-10-20')
        self.assertEqual(res.status_code, 200)
        data = res.data

        comp = data['comparison']
        self.assertEqual(comp['previous_period']['total_income'], '6000000.00')
        self.assertEqual(comp['previous_period']['total_expense'], '1000000.00')
        # Income change: (8,000,000 - 6,000,000) / 6,000,000 * 100 = 33.33%
        self.assertEqual(comp['income_change_percentage'], 33.33)
        # Expense change: (2,500,000 - 1,000,000) / 1,000,000 * 100 = 150.0%
        self.assertEqual(comp['expense_change_percentage'], 150.0)

        # Largest spending increases
        increases = data['largest_spending_increases']
        self.assertTrue(len(increases) >= 1)
        self.assertEqual(increases[0]['label_name'], 'Food')
        self.assertEqual(increases[0]['increase_amount'], '1500000.00')
        self.assertEqual(increases[0]['increase_percentage'], 150.0)

    def test_report_export_csv_and_injection_prevention(self):
        # Create normal transaction and potentially malicious formula title
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=120000, date=datetime.date(2026, 10, 2),
            title='=cmd|/c calc!A0'
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.salary_label, type='EARNING', value=5000000, date=datetime.date(2026, 10, 1),
            title='Lương tháng 10'
        )

        res = self.client.get('/api/reports/export/csv/?start_date=2026-10-01&end_date=2026-10-31')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], 'text/csv; charset=utf-8-sig')
        self.assertIn('attachment; filename=', res['Content-Disposition'])

        content = res.content.decode('utf-8-sig')
        # Check metadata in header
        self.assertIn('BÁO CÁO TÀI CHÍNH MONEYMANAGER', content)
        self.assertIn('01/10/2026 đến 31/10/2026', content)
        # Formula injection check: '=cmd...' MUST be sanitized to ''=cmd...'
        self.assertIn("'=cmd|/c calc!A0", content)
        self.assertIn('Lương tháng 10', content)

    def test_report_user_isolation(self):
        other_user = self.factory.create_user()
        other_wallet = other_user.get_wallet()
        other_label = CustomLabel.objects.create(wallet=other_wallet, name='SecretOther', color='#000000')
        self.factory.create_transaction_filled(
            other_wallet, label=other_label, type='EXPENSE', value=99999999, date=datetime.date(2026, 10, 5)
        )

        res = self.client.get('/api/reports/summary/?start_date=2026-10-01&end_date=2026-10-31')
        self.assertEqual(res.status_code, 200)
        # Other user's expense of 99,999,999 must NOT appear in this user's summary
        self.assertNotEqual(res.data['summary']['total_expense'], '99999999.00')
        breakdown_names = [item['label_name'] for item in res.data['category_breakdown']]
        self.assertNotIn('SecretOther', breakdown_names)

    def test_report_export_excel(self):
        self.factory.create_transaction_filled(
            self.wallet, label=self.salary_label, type='EARNING', value=5000000, date=datetime.date(2026, 10, 1),
            title='Lương tháng 10'
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.food_label, type='EXPENSE', value=120000, date=datetime.date(2026, 10, 2),
            title='Ăn uống trưa'
        )

        res = self.client.get('/api/reports/export/excel/?start_date=2026-10-01&end_date=2026-10-31')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', res['Content-Disposition'])

        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws['A2'].value, 'BÁO CÁO TÀI CHÍNH MONEYMANAGER')
        # Check table headers in row 8
        self.assertEqual(ws['A8'].value, 'Mã GD')
        self.assertEqual(ws['B8'].value, 'Ngày')
        self.assertEqual(ws['C8'].value, 'Tiêu đề')
        self.assertEqual(ws['D8'].value, 'Loại')
        self.assertEqual(ws['E8'].value, 'Danh mục')
        self.assertEqual(ws['F8'].value, 'Số tiền (VND)')
        # Check data exists in columns
        self.assertTrue(ws['A9'].value is not None)
        self.assertTrue(ws['C9'].value in ['Lương tháng 10', 'Ăn uống trưa'])

    def test_import_preview_and_import_invoices(self):
        import os
        from django.core.files.uploadedfile import SimpleUploadedFile
        sample_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../sample_invoices/bien_lai_vietcombank.png'))
        if os.path.exists(sample_path):
            with open(sample_path, 'rb') as f:
                upload_file = SimpleUploadedFile('bien_lai_vietcombank.png', f.read(), content_type='image/png')

            # 1. Test Preview API
            res_preview = self.client.post(
                '/api/import/preview/?institution=Ngân hàng Việt Nam (Chung / VietQR)',
                {'files[]': upload_file},
                format='multipart'
            )
            self.assertEqual(res_preview.status_code, 200)
            self.assertTrue(res_preview.data['success'])
            self.assertEqual(len(res_preview.data['results']), 1)
            first_result = res_preview.data['results'][0]
            self.assertIn(first_result['value'], [850000.0, 880000.0])
            self.assertTrue('850.000' in first_result['formatted_value'] or '880.000' in first_result['formatted_value'])
            self.assertEqual(first_result['formatted_date'], '11/10/2026')

            # 2. Test Import API
            with open(sample_path, 'rb') as f:
                upload_file_2 = SimpleUploadedFile('bien_lai_vietcombank.png', f.read(), content_type='image/png')
            res_import = self.client.post(
                '/api/import/?type=EXPENSE&updateWallet=1',
                {'files[]': upload_file_2},
                format='multipart'
            )
            self.assertEqual(res_import.status_code, 200)
            self.assertTrue(res_import.data['success'])
            self.assertTrue(len(res_import.data['imported']) >= 1)
            self.assertIn(res_import.data['imported'][0]['value'], [850000.0, 880000.0])


class AuthenticationAndAuthorizationTestCase(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.factory = TransactionFactory()

    def test_user_registration_success(self):
        payload = {
            'firstName': 'Nguyen',
            'lastName': 'Van A',
            'email': 'nguyenvana@example.com',
            'password': 'SecurePassword123!',
        }
        response = self.client.post('/api/register/', payload, format='json')
        self.assertEqual(response.status_code, 201)

        user = HubUser.objects.filter(email='nguyenvana@example.com').first()
        self.assertIsNotNone(user)
        self.assertEqual(user.first_name, 'Nguyen')
        self.assertEqual(user.last_name, 'Van A')
        self.assertTrue(user.check_password('SecurePassword123!'))

        wallet = user.get_wallet()
        self.assertIsNotNone(wallet)
        self.assertEqual(wallet.current_amount, Decimal('0.00'))

    def test_user_registration_duplicate_email(self):
        self.factory.create_user(email='existing@example.com')
        payload = {
            'firstName': 'Duplicate',
            'lastName': 'User',
            'email': 'existing@example.com',
            'password': 'Password123!',
        }
        response = self.client.post('/api/register/', payload, format='json')
        self.assertEqual(response.status_code, 400)

    def test_user_registration_invalid_inputs(self):
        # Empty email
        res1 = self.client.post('/api/register/', {'email': '', 'password': 'Pass'}, format='json')
        self.assertEqual(res1.status_code, 400)

        # Invalid email format
        res2 = self.client.post('/api/register/', {'email': 'invalid-email', 'password': 'Pass'}, format='json')
        self.assertEqual(res2.status_code, 400)

        # Empty password
        res3 = self.client.post('/api/register/', {'email': 'valid@example.com', 'password': ''}, format='json')
        self.assertEqual(res3.status_code, 400)

        # Non-dictionary payload
        res4 = self.client.post('/api/register/', 'not-a-dict', format='json')
        self.assertEqual(res4.status_code, 400)

    def test_token_obtain_and_refresh_flow(self):
        user = HubUser.objects.create_user(
            first_name='Test',
            last_name='User',
            email='jwtflow@example.com',
            password='MyPassword123!'
        )

        # 1. Login with valid credentials
        res_login = self.client.post('/api/token/', {
            'email': 'jwtflow@example.com',
            'password': 'MyPassword123!',
        }, format='json')
        self.assertEqual(res_login.status_code, 200)
        self.assertIn('access', res_login.data)
        self.assertIn('refresh_token', res_login.cookies)
        refresh_cookie = res_login.cookies['refresh_token']
        self.assertTrue(refresh_cookie['httponly'])

        # 2. Refresh token using cookie
        self.client.cookies['refresh_token'] = refresh_cookie.value
        res_refresh = self.client.post('/api/token/refresh/')
        self.assertEqual(res_refresh.status_code, 200)
        self.assertIn('access_token', res_refresh.data)

    def test_token_obtain_invalid_credentials(self):
        HubUser.objects.create_user(
            first_name='Test',
            last_name='User',
            email='wrongpwd@example.com',
            password='CorrectPassword!'
        )
        res = self.client.post('/api/token/', {
            'email': 'wrongpwd@example.com',
            'password': 'WrongPassword!',
        }, format='json')
        self.assertEqual(res.status_code, 401)

    def test_token_refresh_without_cookie_fails(self):
        self.client.cookies.clear()
        res = self.client.post('/api/token/refresh/')
        self.assertEqual(res.status_code, 403)

    def test_unauthenticated_requests_are_denied(self):
        anonymous_client = APIClient()
        endpoints = [
            ('GET', '/api/wallet/'),
            ('GET', '/api/transactions/'),
            ('POST', '/api/transaction/'),
            ('GET', '/api/budgets/'),
            ('GET', '/api/budgets/summary/'),
            ('GET', '/api/savingplans/'),
            ('GET', '/api/reports/summary/'),
            ('GET', '/api/reports/transactions/'),
            ('POST', '/api/import/'),
            ('POST', '/api/import/preview/'),
        ]
        for method, url in endpoints:
            if method == 'GET':
                resp = anonymous_client.get(url)
            else:
                resp = anonymous_client.post(url, {}, format='json')
            self.assertEqual(
                resp.status_code,
                401,
                f"Expected 401 Unauthorized for {method} {url}, got {resp.status_code}"
            )


class ObjectOwnershipAndAccessControlTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user_a = self.factory.create_user(email='usera@example.com')
        self.wallet_a = self.user_a.get_wallet()
        self.wallet_a.set_balance(1000)

        self.user_b = self.factory.create_user(email='userb@example.com')
        self.wallet_b = self.user_b.get_wallet()
        self.wallet_b.set_balance(500)

        self.client_a = APIClient()
        self.client_a.force_authenticate(self.user_a)

    def test_user_a_cannot_view_or_update_user_b_transaction(self):
        tx_b = self.factory.create_transaction_filled(
            self.wallet_b, title='Secret B', value=200, type='EXPENSE', update_wallet=True
        )

        # User A tries to edit User B's transaction
        res_put = self.client_a.put(f'/api/transaction/{tx_b.pk}', {
            'title': 'Hacked',
            'value': 10,
            'type': 'EXPENSE',
            'date': '2026-10-01'
        }, format='json')
        self.assertEqual(res_put.status_code, 404)

        # User A tries to delete User B's transaction
        res_del = self.client_a.delete(f'/api/transaction/{tx_b.pk}')
        self.assertEqual(res_del.status_code, 404)

        # Verify wallet B was not modified
        self.wallet_b.refresh_from_db()
        self.assertEqual(self.wallet_b.current_amount, Decimal('300.00'))

    def test_user_a_cannot_assign_user_b_label_to_transaction(self):
        label_b = CustomLabel.objects.create(wallet=self.wallet_b, name='B Category', color='#123456')

        res = self.client_a.post('/api/transaction/', {
            'title': 'My Transaction',
            'value': 100,
            'type': 'EXPENSE',
            'date': '2026-10-01',
            'label': {'id': label_b.pk}
        }, format='json')
        self.assertEqual(res.status_code, 400)

    def test_user_a_cannot_delete_user_b_custom_label(self):
        label_b = CustomLabel.objects.create(wallet=self.wallet_b, name='Private Label', color='#abcdef')
        res = self.client_a.delete(f'/api/labels/{label_b.pk}')
        self.assertEqual(res.status_code, 404)
        self.assertTrue(CustomLabel.objects.filter(pk=label_b.pk).exists())

    def test_user_a_cannot_access_or_modify_user_b_budget(self):
        label_b = CustomLabel.objects.create(wallet=self.wallet_b, name='Budget B', color='#223344')
        budget_b = CategoryBudget.objects.create(
            wallet=self.wallet_b, label=label_b, amount=Decimal('500.00'), month=10, year=2026
        )

        # User A attempts GET on User B's budget
        res_get = self.client_a.get(f'/api/budgets/{budget_b.pk}')
        self.assertEqual(res_get.status_code, 404)

        # User A attempts PUT
        res_put = self.client_a.put(f'/api/budgets/{budget_b.pk}', {'amount': 100}, format='json')
        self.assertEqual(res_put.status_code, 404)

        # User A attempts DELETE
        res_del = self.client_a.delete(f'/api/budgets/{budget_b.pk}')
        self.assertEqual(res_del.status_code, 404)

        # User A attempts to create budget with User B's category
        res_create = self.client_a.post('/api/budgets/', {
            'label_id': label_b.pk,
            'amount': 300,
            'month': 10,
            'year': 2026,
        }, format='json')
        self.assertEqual(res_create.status_code, 400)

    def test_user_a_cannot_access_or_contribute_to_user_b_saving_plan(self):
        plan_b = SavingPlan.objects.create(
            wallet=self.wallet_b, title='Car Fund', amount=Decimal('10000.00')
        )

        # User A attempts GET
        res_get = self.client_a.get(f'/api/savingplans/{plan_b.pk}')
        self.assertEqual(res_get.status_code, 404)

        # User A attempts PUT
        res_put = self.client_a.put(f'/api/savingplans/{plan_b.pk}', {'title': 'Stolen Plan'}, format='json')
        self.assertEqual(res_put.status_code, 404)

        # User A attempts DELETE
        res_del = self.client_a.delete(f'/api/savingplans/{plan_b.pk}')
        self.assertEqual(res_del.status_code, 404)

        # User A attempts to add a contribution to User B's plan
        res_contrib = self.client_a.post(f'/api/savingplans/{plan_b.pk}/contributions/', {
            'amount': 500,
            'date': '2026-10-15',
        }, format='json')
        self.assertEqual(res_contrib.status_code, 404)
        self.assertEqual(plan_b.contributions.count(), 0)

    def test_reports_wallet_isolation(self):
        # User A attempting to query User B's wallet in reports summary returns 403 Forbidden
        res = self.client_a.get(f'/api/reports/summary/?wallet_id={self.wallet_b.pk}')
        self.assertEqual(res.status_code, 403)


class InvoiceImportValidationTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user(email='importsec@example.com')
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_import_without_files_returns_400(self):
        res1 = self.client.post('/api/import/')
        self.assertEqual(res1.status_code, 400)

        res2 = self.client.post('/api/import/preview/')
        self.assertEqual(res2.status_code, 400)

    def test_import_unsupported_file_extension_returns_400(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        fake_exe = SimpleUploadedFile('exploit.exe', b'MZExecutableBinaryData', content_type='application/x-msdownload')

        res_import = self.client.post('/api/import/', {'files[]': fake_exe}, format='multipart')
        self.assertEqual(res_import.status_code, 400)

        fake_sh = SimpleUploadedFile('script.sh', b'#!/bin/bash\necho 1', content_type='text/x-shellscript')
        res_preview = self.client.post('/api/import/preview/', {'files[]': fake_sh}, format='multipart')
        self.assertEqual(res_preview.status_code, 400)

    def test_import_oversized_file_returns_400(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        # 11 MB simulated file
        oversized = SimpleUploadedFile('big_invoice.pdf', b'a' * (11 * 1024 * 1024), content_type='application/pdf')
        res = self.client.post('/api/import/', {'files[]': oversized}, format='multipart')
        self.assertEqual(res.status_code, 400)

    def test_import_corrupted_file_handles_gracefully(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        corrupted = SimpleUploadedFile('corrupt.pdf', b'NOT_A_VALID_PDF_HEADER', content_type='application/pdf')
        res = self.client.post('/api/import/', {'files[]': corrupted}, format='multipart')
        # Should return user-friendly error response, not unhandled 500
        self.assertIn(res.status_code, [400, 422])


class BoundaryAndValidationTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user(email='boundary@example.com')
        self.wallet = self.user.get_wallet()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_transaction_negative_and_zero_amount(self):
        # Value = 0
        res0 = self.client.post('/api/transaction/', {
            'title': 'Zero Test', 'value': 0, 'type': 'EXPENSE', 'date': '2026-10-15'
        }, format='json')
        self.assertEqual(res0.status_code, 400)

        # Value < 0
        res_neg = self.client.post('/api/transaction/', {
            'title': 'Negative Test', 'value': -500, 'type': 'EXPENSE', 'date': '2026-10-15'
        }, format='json')
        self.assertEqual(res_neg.status_code, 400)

        # Non-numeric
        res_str = self.client.post('/api/transaction/', {
            'title': 'String Test', 'value': 'abc', 'type': 'EXPENSE', 'date': '2026-10-15'
        }, format='json')
        self.assertEqual(res_str.status_code, 400)

    def test_transaction_invalid_dates_and_types(self):
        # Invalid date
        res_date = self.client.post('/api/transaction/', {
            'title': 'Bad Date', 'value': 100, 'type': 'EXPENSE', 'date': 'invalid-date'
        }, format='json')
        self.assertEqual(res_date.status_code, 400)

        # Invalid type
        res_type = self.client.post('/api/transaction/', {
            'title': 'Bad Type', 'value': 100, 'type': 'TRANSFER', 'date': '2026-10-15'
        }, format='json')
        self.assertEqual(res_type.status_code, 400)

    def test_wallet_balance_invalid_values(self):
        res1 = self.client.put('/api/wallet/', 'not-a-number', format='json')
        self.assertEqual(res1.status_code, 422)

        res2 = self.client.put('/api/wallet/', True, format='json')
        self.assertEqual(res2.status_code, 422)

    def test_budget_boundary_month_and_year(self):
        label = CustomLabel.objects.create(wallet=self.wallet, name='Food', color='#FF0000')

        # Month < 1
        res1 = self.client.post('/api/budgets/', {
            'label_id': label.pk, 'amount': 100, 'month': 0, 'year': 2026
        }, format='json')
        self.assertEqual(res1.status_code, 400)

        # Month > 12
        res2 = self.client.post('/api/budgets/', {
            'label_id': label.pk, 'amount': 100, 'month': 13, 'year': 2026
        }, format='json')
        self.assertEqual(res2.status_code, 400)

        # Year < 2000
        res3 = self.client.post('/api/budgets/', {
            'label_id': label.pk, 'amount': 100, 'month': 5, 'year': 1999
        }, format='json')
        self.assertEqual(res3.status_code, 400)

        # Year > 2100
        res4 = self.client.post('/api/budgets/', {
            'label_id': label.pk, 'amount': 100, 'month': 5, 'year': 2101
        }, format='json')
        self.assertEqual(res4.status_code, 400)

        # Amount <= 0
        res5 = self.client.post('/api/budgets/', {
            'label_id': label.pk, 'amount': -50, 'month': 5, 'year': 2026
        }, format='json')
        self.assertEqual(res5.status_code, 400)

    def test_saving_plan_boundary_amounts(self):
        # Negative target amount
        res1 = self.client.post('/api/savingplans/', {
            'title': 'Test Plan', 'amount': -100
        }, format='json')
        self.assertEqual(res1.status_code, 400)

        # Zero target amount
        res2 = self.client.post('/api/savingplans/', {
            'title': 'Test Plan', 'amount': 0
        }, format='json')
        self.assertEqual(res2.status_code, 400)

        # Negative contribution amount
        plan = SavingPlan.objects.create(wallet=self.wallet, title='Goal', amount=Decimal('1000.00'))
        res3 = self.client.post(f'/api/savingplans/{plan.pk}/contributions/', {
            'amount': -50
        }, format='json')
        self.assertEqual(res3.status_code, 400)

    def test_nonexistent_resource_endpoints_return_404(self):
        self.assertEqual(self.client.get('/api/budget/999999').status_code, 404)
        self.assertEqual(self.client.delete('/api/transaction/999999').status_code, 404)
        self.assertEqual(self.client.get('/api/savingplan/999999').status_code, 404)
        self.assertEqual(self.client.delete('/api/labels/999999').status_code, 404)


class BudgetWarningAndRebalanceTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

        # Other user for testing authorization
        self.other_user = self.factory.create_user()
        self.other_wallet = self.other_user.get_wallet()
        self.other_client = APIClient()
        self.other_client.force_authenticate(self.other_user)

        self.label_a = CustomLabel.objects.create(wallet=self.wallet, name='Ăn uống', color='#EF4444')
        self.label_b = CustomLabel.objects.create(wallet=self.wallet, name='Giải trí & Cà phê', color='#EC4899')
        self.label_c = CustomLabel.objects.create(wallet=self.wallet, name='Tiền nhà', color='#3B82F6')
        self.other_label = CustomLabel.objects.create(wallet=self.other_wallet, name='Other Label', color='#10B981')

    def test_check_impact_overspent_exceeded(self):
        """
        Tiêu chí nghiệm thu:
        Ngân sách 3.000.000đ, đã chi 2.800.000đ, kiểm tra thêm 500.000đ:
        overspent_amount = 300.000, status = exceeded.
        """
        budget = CategoryBudget.objects.create(
            wallet=self.wallet,
            label=self.label_a,
            amount=Decimal('3000000.00'),
            month=5,
            year=2026,
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_a, type='EXPENSE', value=2800000, date=datetime.date(2026, 5, 10)
        )

        res = self.client.post('/api/budgets/check-impact/', {
            'label_id': self.label_a.pk,
            'value': 500000,
            'date': '2026-05-15',
        }, format='json')

        self.assertEqual(res.status_code, 200)
        data = res.data
        self.assertTrue(data['has_budget'])
        self.assertEqual(data['budget_amount'], '3000000.00')
        self.assertEqual(data['spent_before'], '2800000.00')
        self.assertEqual(data['spent_after'], '3300000.00')
        self.assertEqual(data['percent_before'], round(2800000.0 / 3000000.0 * 100, 2))
        self.assertEqual(data['percent_after'], 110.0)
        self.assertEqual(data['overspent_amount'], '300000.00')
        self.assertEqual(data['status'], 'exceeded')

    def test_check_impact_no_budget(self):
        """
        Tiêu chí nghiệm thu:
        Danh mục không có ngân sách: không có cảnh báo, không lỗi.
        """
        res = self.client.post('/api/budgets/check-impact/', {
            'label_id': self.label_a.pk,
            'value': 500000,
            'date': '2026-06-15',
        }, format='json')

        self.assertEqual(res.status_code, 200)
        data = res.data
        self.assertFalse(data['has_budget'])
        self.assertEqual(data['overspent_amount'], '0.00')
        self.assertEqual(data['status'], 'normal')

    def test_check_impact_approaching(self):
        """
        Kiểm tra ngưỡng tiệm cận >= 80% (approaching).
        """
        budget = CategoryBudget.objects.create(
            wallet=self.wallet,
            label=self.label_a,
            amount=Decimal('1000000.00'),
            month=5,
            year=2026,
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_a, type='EXPENSE', value=700000, date=datetime.date(2026, 5, 10)
        )

        res = self.client.post('/api/budgets/check-impact/', {
            'label_id': self.label_a.pk,
            'value': 150000,
            'date': '2026-05-15',
        }, format='json')

        self.assertEqual(res.status_code, 200)
        data = res.data
        self.assertTrue(data['has_budget'])
        self.assertEqual(data['percent_after'], 85.0)
        self.assertEqual(data['overspent_amount'], '0.00')
        self.assertEqual(data['status'], 'approaching')

    def test_rebalance_success(self):
        """
        Tiêu chí nghiệm thu:
        Rebalance 300.000đ từ B sang A: B giảm, A tăng đúng 300.000, tổng ngân sách không đổi.
        """
        budget_a = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_a, amount=Decimal('1000000.00'), month=5, year=2026
        )
        budget_b = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_b, amount=Decimal('2000000.00'), month=5, year=2026
        )
        # B đã chi 500.000đ, vẫn còn 1.500.000đ khả dụng
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_b, type='EXPENSE', value=500000, date=datetime.date(2026, 5, 2)
        )

        total_before = budget_a.amount + budget_b.amount

        res = self.client.post('/api/budgets/rebalance/', {
            'from_label_id': self.label_b.pk,
            'to_label_id': self.label_a.pk,
            'amount': 300000,
            'month': 5,
            'year': 2026,
        }, format='json')

        self.assertEqual(res.status_code, 200)

        budget_a.refresh_from_db()
        budget_b.refresh_from_db()

        self.assertEqual(budget_b.amount, Decimal('1700000.00'))
        self.assertEqual(budget_a.amount, Decimal('1300000.00'))
        self.assertEqual(budget_a.amount + budget_b.amount, total_before)

    def test_rebalance_atomic_rollback_on_failure(self):
        """
        Tiêu chí nghiệm thu:
        Nếu một bước cập nhật lỗi thì cả hai ngân sách giữ nguyên (atomic).
        Ngân sách nguồn sau khi giảm không được nhỏ hơn số đã chi.
        """
        budget_a = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_a, amount=Decimal('1000000.00'), month=5, year=2026
        )
        budget_b = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_b, amount=Decimal('1000000.00'), month=5, year=2026
        )
        # B đã chi 850.000đ, chỉ còn 150.000đ khả dụng
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_b, type='EXPENSE', value=850000, date=datetime.date(2026, 5, 2)
        )

        # Thử rebalance 300.000đ (vượt quá số dư khả dụng của B)
        res = self.client.post('/api/budgets/rebalance/', {
            'from_label_id': self.label_b.pk,
            'to_label_id': self.label_a.pk,
            'amount': 300000,
            'month': 5,
            'year': 2026,
        }, format='json')

        self.assertEqual(res.status_code, 400)

        # Cả 2 ngân sách giữ nguyên
        budget_a.refresh_from_db()
        budget_b.refresh_from_db()
        self.assertEqual(budget_a.amount, Decimal('1000000.00'))
        self.assertEqual(budget_b.amount, Decimal('1000000.00'))

    def test_rebalance_atomic_rollback_on_destination_save_failure(self):
        """
        Tiêu chí nghiệm thu A2:
        Atomic rollback kiểm chứng thực sự: save() của ngân sách đích ném exception sau khi
        ngân sách nguồn đã save() thành công.
        Kiểm tra từ DB (refresh_from_db) rằng cả hai ngân sách vẫn giữ nguyên số tiền ban đầu.
        """
        budget_a = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_a, amount=Decimal('1000000.00'), month=5, year=2026
        )
        budget_b = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_b, amount=Decimal('2000000.00'), month=5, year=2026
        )
        # B đã chi 500.000đ, còn 1.500.000đ khả dụng (thỏa mãn điều kiện validate)
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_b, type='EXPENSE', value=500000, date=datetime.date(2026, 5, 2)
        )

        original_save = CategoryBudget.save

        def faulty_save(instance, *args, **kwargs):
            # Nếu là ngân sách đích (to_budget là budget_a), ném exception sau khi ngân sách nguồn (budget_b) đã save thành công
            if instance.pk == budget_a.pk:
                raise RuntimeError("Simulated failure when saving destination budget")
            return original_save(instance, *args, **kwargs)

        with mock.patch.object(CategoryBudget, 'save', autospec=True, side_effect=faulty_save):
            res = self.client.post('/api/budgets/rebalance/', {
                'from_label_id': self.label_b.pk,
                'to_label_id': self.label_a.pk,
                'amount': 300000,
                'month': 5,
                'year': 2026,
            }, format='json')

        self.assertEqual(res.status_code, 500)
        self.assertIn('hoàn tác', res.json().get('message', ''))

        # Cả hai ngân sách phải được rollback về số tiền ban đầu
        budget_a.refresh_from_db()
        budget_b.refresh_from_db()
        self.assertEqual(budget_a.amount, Decimal('1000000.00'))
        self.assertEqual(budget_b.amount, Decimal('2000000.00'))

    def test_object_ownership_authorization(self):
        """
        Tiêu chí nghiệm thu:
        Người dùng khác gọi với label/budget của người khác: trả 403 hoặc 404.
        """
        budget_a = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_a, amount=Decimal('1000000.00'), month=5, year=2026
        )

        # Other user calls check-impact with user's label
        res1 = self.other_client.post('/api/budgets/check-impact/', {
            'label_id': self.label_a.pk,
            'value': 100000,
            'date': '2026-05-15',
        }, format='json')
        self.assertEqual(res1.status_code, 404)

        # Other user calls rebalance with user's label
        res2 = self.other_client.post('/api/budgets/rebalance/', {
            'from_label_id': self.label_a.pk,
            'to_label_id': self.other_label.pk,
            'amount': 100000,
            'month': 5,
            'year': 2026,
        }, format='json')
        self.assertEqual(res2.status_code, 404)

        # Other user calls rebalance-suggestions with user's label
        res3 = self.other_client.get(f'/api/budgets/rebalance-suggestions/?label_id={self.label_a.pk}&month=5&year=2026')
        self.assertEqual(res3.status_code, 404)

    def test_rebalance_suggestions_filtering_and_priority(self):
        """
        Tiêu chí nghiệm thu:
        - Danh mục đã chi >= 60% hoặc mới đầu tháng (chưa qua 50% số ngày) không xuất hiện trong gợi ý.
        - Đủ điều kiện: có ngân sách, đã chi < 60%, đã qua hơn 50% số ngày, số dư còn lại >= thâm hụt.
        - Ưu tiên danh mục không thiết yếu trước.
        """
        # Test 1: Tháng trong tương lai (0% số ngày đã qua) -> trả về []
        future_year = 2099
        CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_b, amount=Decimal('1000000.00'), month=5, year=future_year
        )
        res_future = self.client.get(f'/api/budgets/rebalance-suggestions/?month=5&year={future_year}')
        self.assertEqual(res_future.status_code, 200)
        self.assertEqual(res_future.data, [])

        # Test 2: Tháng trong quá khứ (ví dụ 10/2020, 100% đã qua):
        past_month, past_year = 10, 2020
        # Target: Ăn uống (ngân sách 1.000.000đ, đã chi 1.300.000đ -> thâm hụt 300.000đ)
        b_target = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_a, amount=Decimal('1000000.00'), month=past_month, year=past_year
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_a, type='EXPENSE', value=1300000, date=datetime.date(past_year, past_month, 5)
        )

        # Nguồn 1: Tiền nhà (thiết yếu, ngân sách 3.000.000đ, đã chi 600.000đ = 20% < 60%, còn lại 2.400.000đ >= 300.000đ) -> Đủ điều kiện
        b_essential = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_c, amount=Decimal('3000000.00'), month=past_month, year=past_year
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_c, type='EXPENSE', value=600000, date=datetime.date(past_year, past_month, 5)
        )

        # Nguồn 2: Giải trí & Cà phê (không thiết yếu, ngân sách 1.000.000đ, đã chi 200.000đ = 20% < 60%, còn lại 800.000đ >= 300.000đ) -> Đủ điều kiện
        b_non_essential = CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_b, amount=Decimal('1000000.00'), month=past_month, year=past_year
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_b, type='EXPENSE', value=200000, date=datetime.date(past_year, past_month, 5)
        )

        # Nguồn 3: Danh mục khác (đã chi 700.000/1.000.000 = 70% >= 60%) -> Bị loại
        label_d = CustomLabel.objects.create(wallet=self.wallet, name='Đi lại', color='#F59E0B')
        CategoryBudget.objects.create(
            wallet=self.wallet, label=label_d, amount=Decimal('1000000.00'), month=past_month, year=past_year
        )
        self.factory.create_transaction_filled(
            self.wallet, label=label_d, type='EXPENSE', value=700000, date=datetime.date(past_year, past_month, 5)
        )

        # Nguồn 4: Danh mục số dư còn lại < thâm hụt (còn lại 100.000đ < 300.000đ thâm hụt) -> Bị loại
        label_e = CustomLabel.objects.create(wallet=self.wallet, name='Mua sắm nhỏ', color='#8B5CF6')
        CategoryBudget.objects.create(
            wallet=self.wallet, label=label_e, amount=Decimal('200000.00'), month=past_month, year=past_year
        )
        self.factory.create_transaction_filled(
            self.wallet, label=label_e, type='EXPENSE', value=100000, date=datetime.date(past_year, past_month, 5)
        )

        res_suggestions = self.client.get(
            f'/api/budgets/rebalance-suggestions/?label_id={self.label_a.pk}&month={past_month}&year={past_year}'
        )
        self.assertEqual(res_suggestions.status_code, 200)
        data = res_suggestions.data
        # Chỉ có 2 danh mục đủ điều kiện
        self.assertEqual(len(data), 2)

        # Danh mục không thiết yếu ("Giải trí & Cà phê") được ưu tiên đứng đầu!
        self.assertEqual(data[0]['label_name'], 'Giải trí & Cà phê')
        self.assertTrue(data[0]['is_non_essential'])
        self.assertEqual(data[0]['amount'], '1000000.00')
        self.assertEqual(data[0]['spent'], '200000.00')
        self.assertEqual(data[0]['available_rebalance'], '800000.00')

        # Danh mục thiết yếu ("Tiền nhà") đứng sau
        self.assertEqual(data[1]['label_name'], 'Tiền nhà')
        self.assertFalse(data[1]['is_non_essential'])
        self.assertEqual(data[1]['amount'], '3000000.00')
        self.assertEqual(data[1]['spent'], '600000.00')
        self.assertEqual(data[1]['available_rebalance'], '2400000.00')

    def test_rebalance_suggestions_unaccented_non_essential_labels(self):
        """
        Tiêu chí B4:
        Từ khóa không thiết yếu nhận diện được cả nhãn không dấu:
        'giai tri', 'mua sam', 'ca phe' -> is_non_essential = True.
        Kiểu dữ liệu tiền tệ trả về chuỗi Decimal theo convention (Tiêu chí B5).
        """
        past_month, past_year = 11, 2020
        # Target: Ăn uống (thâm hụt 100.000đ)
        CategoryBudget.objects.create(
            wallet=self.wallet, label=self.label_a, amount=Decimal('500000.00'), month=past_month, year=past_year
        )
        self.factory.create_transaction_filled(
            self.wallet, label=self.label_a, type='EXPENSE', value=600000, date=datetime.date(past_year, past_month, 5)
        )

        # Nguồn không dấu: "giai tri cuoi tuan", "mua sam online", "ca phe sang" và thiết yếu "tien dien nuoc"
        lbl_giai_tri = CustomLabel.objects.create(wallet=self.wallet, name='giai tri cuoi tuan', color='#111111')
        lbl_mua_sam = CustomLabel.objects.create(wallet=self.wallet, name='mua sam online', color='#222222')
        lbl_ca_phe = CustomLabel.objects.create(wallet=self.wallet, name='ca phe sang', color='#333333')
        lbl_dien_nuoc = CustomLabel.objects.create(wallet=self.wallet, name='tien dien nuoc', color='#444444')

        for lbl in [lbl_giai_tri, lbl_mua_sam, lbl_ca_phe, lbl_dien_nuoc]:
            CategoryBudget.objects.create(
                wallet=self.wallet, label=lbl, amount=Decimal('1000000.00'), month=past_month, year=past_year
            )

        res = self.client.get(
            f'/api/budgets/rebalance-suggestions/?label_id={self.label_a.pk}&month={past_month}&year={past_year}'
        )
        self.assertEqual(res.status_code, 200)
        data = res.data

        # Tìm các item theo tên
        item_giai_tri = next(i for i in data if i['label_name'] == 'giai tri cuoi tuan')
        item_mua_sam = next(i for i in data if i['label_name'] == 'mua sam online')
        item_ca_phe = next(i for i in data if i['label_name'] == 'ca phe sang')
        item_dien_nuoc = next(i for i in data if i['label_name'] == 'tien dien nuoc')

        self.assertTrue(item_giai_tri['is_non_essential'])
        self.assertTrue(item_mua_sam['is_non_essential'])
        self.assertTrue(item_ca_phe['is_non_essential'])
        self.assertFalse(item_dien_nuoc['is_non_essential'])

        # Kiểm tra kiểu dữ liệu chuỗi Decimal (tiêu chí B5)
        self.assertEqual(item_giai_tri['amount'], '1000000.00')
        self.assertEqual(item_giai_tri['available_rebalance'], '1000000.00')


class DebtManagementTestCase(TestCase):
    def setUp(self):
        self.factory = TransactionFactory()
        self.user = self.factory.create_user()
        self.wallet = self.user.get_wallet()
        self.wallet.set_balance(Decimal('10000000.00'))  # 10.000.000đ

        self.other_user = self.factory.create_user()
        self.other_wallet = self.other_user.get_wallet()
        self.other_wallet.set_balance(Decimal('5000000.00'))

        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_01_loan_given_updates_wallet(self):
        """Cho vay 5.000.000đ (update_wallet bật): ví giảm 5.000.000đ."""
        initial_balance = self.wallet.get_current_amount()
        res = self.client.post('/api/debts/', {
            'person_name': 'Nguyen Van A',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': True,
            'start_date': '2026-10-01',
            'due_date': '2026-11-01',
        }, format='json')
        self.assertEqual(res.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance - Decimal('5000000.00'))
        debt_id = res.data['id']
        debt = DebtLoan.objects.get(pk=debt_id)
        self.assertEqual(debt.get_remaining_amount(), Decimal('5000000.00'))
        self.assertEqual(debt.get_computed_status(), 'ACTIVE')

    def test_02_loan_collection_and_full_payment(self):
        """Thu hồi 2.000.000đ: dư nợ 3.000.000, PARTIALLY_PAID, ví tăng 2.000.000. Thu hồi tiếp 3.000.000: FULLY_PAID."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Tran Van B',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': True,
        }, user_pk=self.user.pk)
        self.wallet.refresh_from_db()
        bal_after_loan = self.wallet.get_current_amount()

        # Thu hồi 2.000.000đ
        res1 = self.client.post(f'/api/debts/{debt.pk}/payments/', {
            'amount': '2000000.00',
            'update_wallet': True,
            'note': 'Thu đợt 1',
        }, format='json')
        self.assertEqual(res1.status_code, 201)
        self.wallet.refresh_from_db()
        debt.refresh_from_db()

        self.assertEqual(debt.get_remaining_amount(), Decimal('3000000.00'))
        self.assertEqual(debt.get_computed_status(), 'PARTIALLY_PAID')
        self.assertEqual(self.wallet.get_current_amount(), bal_after_loan + Decimal('2000000.00'))

        # Thu hồi tiếp 3.000.000đ
        res2 = self.client.post(f'/api/debts/{debt.pk}/payments/', {
            'amount': '3000000.00',
            'update_wallet': True,
            'note': 'Thu đợt 2',
        }, format='json')
        self.assertEqual(res2.status_code, 201)
        self.wallet.refresh_from_db()
        debt.refresh_from_db()

        self.assertEqual(debt.get_remaining_amount(), Decimal('0.00'))
        self.assertEqual(debt.get_computed_status(), 'FULLY_PAID')
        self.assertEqual(self.wallet.get_current_amount(), bal_after_loan + Decimal('5000000.00'))

    def test_03_payment_exceeding_remaining_amount_rejected(self):
        """Thanh toán vượt dư nợ: bị từ chối (400/422), ví không đổi."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Le Van C',
            'type': 'LOAN_GIVEN',
            'principal_amount': '1000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)
        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        res = self.client.post(f'/api/debts/{debt.pk}/payments/', {
            'amount': '1500000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res.status_code, 400)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), bal_before)
        self.assertEqual(debt.payments.count(), 0)

    def test_04_overdue_status_and_days_calculation(self):
        """Quá hạn mà còn nợ: trả OVERDUE và số ngày quá hạn đúng."""
        past_date = datetime.date.today() - datetime.timedelta(days=10)
        start_date = past_date - datetime.timedelta(days=30)
        debt = DebtLoan.create_from_json({
            'person_name': 'Pham Van D',
            'type': 'DEBT_BORROWED',
            'principal_amount': '4000000.00',
            'start_date': start_date.isoformat(),
            'due_date': past_date.isoformat(),
            'update_wallet': False,
        }, user_pk=self.user.pk)

        res = self.client.get(f'/api/debts/{debt.pk}')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['status'], 'OVERDUE')
        self.assertEqual(res.data['overdue_days'], 10)

    def test_05_installment_schedule_calculation(self):
        """Trả góp 24.000.000đ / 12 tháng: 2.000.000đ mỗi kỳ, 12 kỳ trong lịch."""
        res = self.client.post('/api/debts/', {
            'person_name': 'Cua hang Dien may X',
            'type': 'INSTALLMENT',
            'principal_amount': '24000000.00',
            'installment_months': 12,
            'start_date': '2026-01-01',
            'update_wallet': False,
        }, format='json')
        self.assertEqual(res.status_code, 201)
        schedule = res.data['installment_schedule']
        self.assertEqual(len(schedule), 12)
        for item in schedule:
            self.assertEqual(item['amount'], '2000000.00')
            self.assertFalse(item['is_paid'])

        debt_pk = res.data['id']
        # Trả 1 kỳ 2.000.000đ
        self.client.post(f'/api/debts/{debt_pk}/payments/', {
            'amount': '2000000.00',
            'update_wallet': False,
        }, format='json')
        res_after = self.client.get(f'/api/debts/{debt_pk}')
        schedule_after = res_after.data['installment_schedule']
        self.assertTrue(schedule_after[0]['is_paid'])
        self.assertFalse(schedule_after[1]['is_paid'])

    def test_06_update_payment_amount_recalculates_wallet_and_remaining(self):
        """Sửa số tiền một lần thanh toán từ 2.000.000 lên 2.500.000: ví và dư nợ được tính lại đúng."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Hoang Van E',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)
        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        # Tạo payment 2.000.000 có update_wallet
        pay = DebtPayment.create_from_json(debt.pk, self.user, {
            'amount': '2000000.00',
            'update_wallet': True,
        })
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), bal_before + Decimal('2000000.00'))

        # Sửa lên 2.500.000đ
        res = self.client.put(f'/api/debts/{debt.pk}/payments/{pay.pk}', {
            'amount': '2500000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res.status_code, 200)
        self.wallet.refresh_from_db()
        debt.refresh_from_db()

        self.assertEqual(self.wallet.get_current_amount(), bal_before + Decimal('2500000.00'))
        self.assertEqual(debt.get_remaining_amount(), Decimal('2500000.00'))

    def test_07_delete_payment_reverses_wallet_and_restores_status(self):
        """Xóa một lần thanh toán: hoàn lại ví, trạng thái quay về đúng."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Vu Van F',
            'type': 'LOAN_GIVEN',
            'principal_amount': '3000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)
        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        pay = DebtPayment.create_from_json(debt.pk, self.user, {
            'amount': '3000000.00',
            'update_wallet': True,
        })
        debt.refresh_from_db()
        self.assertEqual(debt.get_computed_status(), 'FULLY_PAID')

        # Xóa thanh toán
        res = self.client.delete(f'/api/debts/{debt.pk}/payments/{pay.pk}')
        self.assertEqual(res.status_code, 200)
        self.wallet.refresh_from_db()
        debt.refresh_from_db()

        self.assertEqual(self.wallet.get_current_amount(), bal_before)
        self.assertEqual(debt.get_remaining_amount(), Decimal('3000000.00'))
        self.assertEqual(debt.get_computed_status(), 'ACTIVE')

    def test_08_error_during_creation_rollbacks_atomically(self):
        """Lỗi giữa chừng thì rollback toàn bộ (ví và payment)."""
        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        # Thử gọi create_from_json với mock lỗi khi lưu
        with mock.patch.object(DebtLoan, 'save', side_effect=RuntimeError('Database explosion')):
            with self.assertRaises(RuntimeError):
                DebtLoan.create_from_json({
                    'person_name': 'Loi Test',
                    'type': 'LOAN_GIVEN',
                    'principal_amount': '1000000.00',
                    'update_wallet': True,
                }, user_pk=self.user.pk)

        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), bal_before)
        self.assertEqual(DebtLoan.objects.filter(person_name='Loi Test').count(), 0)

    def test_09_object_level_authorization_isolation(self):
        """Người dùng khác không xem/sửa được hợp đồng của người khác."""
        debt_user1 = DebtLoan.create_from_json({
            'person_name': 'Bi Mat User 1',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # Chuyển client sang user 2
        client2 = APIClient()
        client2.force_authenticate(user=self.other_user)

        # Xem danh sách -> không có debt của user 1
        res_list = client2.get('/api/debts/')
        self.assertEqual(res_list.status_code, 200)
        self.assertEqual(len(res_list.data), 0)

        # Truy cập trực tiếp -> 404
        res_detail = client2.get(f'/api/debts/{debt_user1.pk}')
        self.assertEqual(res_detail.status_code, 404)

        # Sửa -> 404
        res_put = client2.put(f'/api/debts/{debt_user1.pk}', {'person_name': 'Hacker'})
        self.assertEqual(res_put.status_code, 404)

        # Xóa -> 404
        res_del = client2.delete(f'/api/debts/{debt_user1.pk}')
        self.assertEqual(res_del.status_code, 404)

        # Thêm payment -> 404
        res_pay = client2.post(f'/api/debts/{debt_user1.pk}/payments/', {'amount': '100000.00'})
        self.assertEqual(res_pay.status_code, 404)

    def test_10_debts_do_not_affect_report_summary(self):
        """Báo cáo thu chi không thay đổi khi có khoản nợ."""
        # Lấy report ban đầu
        res_initial = self.client.get('/api/reports/summary/')
        self.assertEqual(res_initial.status_code, 200)
        income_before = res_initial.data['summary']['total_income']
        expense_before = res_initial.data['summary']['total_expense']

        # Tạo các khoản nợ có update_wallet
        debt = DebtLoan.create_from_json({
            'person_name': 'Doi Tac Vay',
            'type': 'LOAN_GIVEN',
            'principal_amount': '2000000.00',
            'update_wallet': True,
        }, user_pk=self.user.pk)

        DebtPayment.create_from_json(debt.pk, self.user, {
            'amount': '1000000.00',
            'update_wallet': True,
        })

        res_after = self.client.get('/api/reports/summary/')
        self.assertEqual(res_after.status_code, 200)
        self.assertEqual(res_after.data['summary']['total_income'], income_before)
        self.assertEqual(res_after.data['summary']['total_expense'], expense_before)


    def test_11_debt_summary_api(self):
        """Kiểm tra API summary: tổng cho vay, tổng nợ, đến hạn trong 7 ngày, quá hạn."""
        today = datetime.date.today()
        # 1. Cho vay 3tr, chưa thu hồi -> người khác nợ mình 3tr
        DebtLoan.create_from_json({
            'person_name': 'Cho vay A',
            'type': 'LOAN_GIVEN',
            'principal_amount': '3000000.00',
            'due_date': (today + datetime.timedelta(days=3)).isoformat(),
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # 2. Đi vay 2tr, đến hạn 5 ngày -> mình nợ người ta 2tr
        DebtLoan.create_from_json({
            'person_name': 'Di vay B',
            'type': 'DEBT_BORROWED',
            'principal_amount': '2000000.00',
            'due_date': (today + datetime.timedelta(days=5)).isoformat(),
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # 3. Quá hạn 1tr
        DebtLoan.create_from_json({
            'person_name': 'Qua han C',
            'type': 'DEBT_BORROWED',
            'principal_amount': '1000000.00',
            'start_date': (today - datetime.timedelta(days=10)).isoformat(),
            'due_date': (today - datetime.timedelta(days=2)).isoformat(),
            'update_wallet': False,
        }, user_pk=self.user.pk)


        res = self.client.get('/api/debts/summary/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['total_loan_given_remaining'], '3000000.00')
        self.assertEqual(res.data['total_debt_borrowed_remaining'], '3000000.00')  # 2tr + 1tr
        self.assertEqual(res.data['due_in_7_days_count'], 2)
        self.assertEqual(res.data['overdue_count'], 1)

    def test_12_borrowed_debt_updates_wallet(self):
        """Đi vay 5.000.000đ (update_wallet bật): ví tăng 5.000.000đ; Trả nợ 2.000.000đ: ví giảm 2.000.000đ."""
        initial_balance = self.wallet.get_current_amount()
        res = self.client.post('/api/debts/', {
            'person_name': 'Chu no X',
            'type': 'DEBT_BORROWED',
            'principal_amount': '5000000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance + Decimal('5000000.00'))

        debt_pk = res.data['id']
        # Trả nợ 2.000.000đ -> ví giảm 2.000.000đ
        res_pay = self.client.post(f'/api/debts/{debt_pk}/payments/', {
            'amount': '2000000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res_pay.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance + Decimal('3000000.00'))

    def test_13_create_debt_without_update_wallet_and_delete(self):
        """Tạo hợp đồng update_wallet=False: ví không đổi, wallet_applied = False. Xóa hợp đồng: ví vẫn không đổi."""
        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        res = self.client.post('/api/debts/', {
            'person_name': 'Doi Tac Khong Vi',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': False,
        }, format='json')
        self.assertEqual(res.status_code, 201)
        self.assertFalse(res.data['wallet_applied'])
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), bal_before)

        debt_pk = res.data['id']
        debt = DebtLoan.objects.get(pk=debt_pk)
        self.assertFalse(debt.wallet_applied)

        # Xóa hợp đồng
        res_del = self.client.delete(f'/api/debts/{debt_pk}')
        self.assertEqual(res_del.status_code, 200)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), bal_before)

    def test_14_delete_loan_given_with_mixed_payments_restores_balance(self):
        """Cho vay 5.000.000đ (update_wallet=True) kèm 2 lần thu hồi (một lần update_wallet=True, một lần False), sau đó xóa hợp đồng: ví về ĐÚNG số dư ban đầu."""
        self.wallet.refresh_from_db()
        initial_balance = self.wallet.get_current_amount()

        # Cho vay 5tr có update_wallet
        res = self.client.post('/api/debts/', {
            'person_name': 'Ban Vay Tien',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res.status_code, 201)
        self.assertTrue(res.data['wallet_applied'])
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance - Decimal('5000000.00'))

        debt_pk = res.data['id']

        # Thu hồi 1: 2.000.000đ có update_wallet
        res_p1 = self.client.post(f'/api/debts/{debt_pk}/payments/', {
            'amount': '2000000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res_p1.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance - Decimal('3000000.00'))

        # Thu hồi 2: 1.000.000đ không update_wallet
        res_p2 = self.client.post(f'/api/debts/{debt_pk}/payments/', {
            'amount': '1000000.00',
            'update_wallet': False,
        }, format='json')
        self.assertEqual(res_p2.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance - Decimal('3000000.00'))

        # Xóa hợp đồng: phải hoàn lại 5tr ban đầu, trừ đi 2tr đã thu hồi qua ví -> ví trở về đúng initial_balance
        res_del = self.client.delete(f'/api/debts/{debt_pk}')
        self.assertEqual(res_del.status_code, 200)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance)

    def test_15_delete_debt_borrowed_with_mixed_payments_restores_balance(self):
        """Tương tự cho DEBT_BORROWED: Đi vay 5.000.000đ, trả 2tr (update_wallet=True) và 1tr (update_wallet=False), xóa hợp đồng -> ví về đúng ban đầu."""
        self.wallet.refresh_from_db()
        initial_balance = self.wallet.get_current_amount()

        # Đi vay 5tr có update_wallet -> ví +5tr
        res = self.client.post('/api/debts/', {
            'person_name': 'Nguoi Cho Vay',
            'type': 'DEBT_BORROWED',
            'principal_amount': '5000000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res.status_code, 201)
        self.assertTrue(res.data['wallet_applied'])
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance + Decimal('5000000.00'))

        debt_pk = res.data['id']

        # Trả 1: 2tr có update_wallet -> ví -2tr
        res_p1 = self.client.post(f'/api/debts/{debt_pk}/payments/', {
            'amount': '2000000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res_p1.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance + Decimal('3000000.00'))

        # Trả 2: 1tr không update_wallet -> ví không đổi
        res_p2 = self.client.post(f'/api/debts/{debt_pk}/payments/', {
            'amount': '1000000.00',
            'update_wallet': False,
        }, format='json')
        self.assertEqual(res_p2.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance + Decimal('3000000.00'))

        # Xóa hợp đồng: hoàn 2tr đã trả, trừ 5tr ban đầu -> ví về đúng initial_balance
        res_del = self.client.delete(f'/api/debts/{debt_pk}')
        self.assertEqual(res_del.status_code, 200)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance)

    def test_16_update_principal_adjusts_wallet_when_wallet_applied(self):
        """Sửa principal từ 5.000.000 lên 6.000.000 khi wallet_applied=True: ví giảm thêm đúng 1.000.000."""
        self.wallet.refresh_from_db()
        initial_balance = self.wallet.get_current_amount()

        # Tạo LOAN_GIVEN 5tr có update_wallet
        res = self.client.post('/api/debts/', {
            'person_name': 'Ban Than',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res.status_code, 201)
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), initial_balance - Decimal('5000000.00'))

        debt_pk = res.data['id']

        # Sửa principal lên 6.000.000đ
        res_put = self.client.put(f'/api/debts/{debt_pk}', {
            'principal_amount': '6000000.00',
        }, format='json')
        self.assertEqual(res_put.status_code, 200)
        self.wallet.refresh_from_db()
        # Ví phải giảm thêm đúng 1.000.000đ -> initial_balance - 6tr
        self.assertEqual(self.wallet.get_current_amount(), initial_balance - Decimal('6000000.00'))

    def test_17_update_principal_below_total_paid_rejected(self):
        """Sửa principal xuống dưới tổng đã trả: bị từ chối, dữ liệu và ví không đổi."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Khach Hang',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # Thanh toán 3.000.000đ
        DebtPayment.create_from_json(debt.pk, self.user, {
            'amount': '3000000.00',
            'update_wallet': False,
        })
        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        # Cố sửa principal xuống 2.000.000đ (< 3.000.000đ đã trả)
        res_put = self.client.put(f'/api/debts/{debt.pk}', {
            'principal_amount': '2000000.00',
        }, format='json')
        self.assertEqual(res_put.status_code, 400)
        debt.refresh_from_db()
        self.assertEqual(debt.principal_amount, Decimal('5000000.00'))
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), bal_before)

    def test_18_changing_type_or_wallet_after_payments_rejected(self):
        """Đổi type hoặc wallet của hợp đồng đã có thanh toán hoặc đã wallet_applied: bị từ chối."""
        # 1. Hợp đồng có thanh toán: cấm đổi type
        debt1 = DebtLoan.create_from_json({
            'person_name': 'Nguoi No 1',
            'type': 'LOAN_GIVEN',
            'principal_amount': '3000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)
        DebtPayment.create_from_json(debt1.pk, self.user, {
            'amount': '1000000.00',
            'update_wallet': False,
        })

        res1 = self.client.put(f'/api/debts/{debt1.pk}', {
            'type': 'DEBT_BORROWED',
        }, format='json')
        self.assertEqual(res1.status_code, 400)

        # 2. Hợp đồng đã wallet_applied: cấm đổi type
        debt2 = DebtLoan.create_from_json({
            'person_name': 'Nguoi No 2',
            'type': 'LOAN_GIVEN',
            'principal_amount': '2000000.00',
            'update_wallet': True,
        }, user_pk=self.user.pk)

        res2 = self.client.put(f'/api/debts/{debt2.pk}', {
            'type': 'DEBT_BORROWED',
        }, format='json')
        self.assertEqual(res2.status_code, 400)

        # 3. Cấm đổi wallet
        res3 = self.client.put(f'/api/debts/{debt1.pk}', {
            'wallet': 9999,
        }, format='json')
        self.assertEqual(res3.status_code, 400)

    def test_19_installment_rounding_10m_3_months(self):
        """10.000.000đ / 3 kỳ: tổng 3 kỳ = đúng 10.000.000, kỳ cuối lớn hơn 2 kỳ đầu 1 đồng."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Mua Dien Thoai',
            'type': 'INSTALLMENT',
            'principal_amount': '10000000.00',
            'installment_months': 3,
            'start_date': '2026-01-01',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        schedule = debt.get_installment_schedule()
        self.assertEqual(len(schedule), 3)
        self.assertEqual(schedule[0]['amount'], '3333333.00')
        self.assertEqual(schedule[1]['amount'], '3333333.00')
        self.assertEqual(schedule[2]['amount'], '3333334.00')

        total = sum(Decimal(item['amount']) for item in schedule)
        self.assertEqual(total, Decimal('10000000.00'))

    def test_20_installment_24m_12_months(self):
        """24.000.000đ / 12 kỳ vẫn cho 12 kỳ x 2.000.000đ."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Mua Laptop',
            'type': 'INSTALLMENT',
            'principal_amount': '24000000.00',
            'installment_months': 12,
            'start_date': '2026-01-01',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        schedule = debt.get_installment_schedule()
        self.assertEqual(len(schedule), 12)
        for item in schedule:
            self.assertEqual(item['amount'], '2000000.00')
        total = sum(Decimal(item['amount']) for item in schedule)
        self.assertEqual(total, Decimal('24000000.00'))

    def test_21_installment_sequential_payment_full_paid(self):
        """Trả đúng theo lịch từng kỳ: sau kỳ cuối, trạng thái = FULLY_PAID và mọi kỳ is_paid = True."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Mua Xe May',
            'type': 'INSTALLMENT',
            'principal_amount': '10000000.00',
            'installment_months': 3,
            'start_date': '2026-01-01',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # Trả kỳ 1: 3.333.333đ
        DebtPayment.create_from_json(debt.pk, self.user, {'amount': '3333333.00'})
        s1 = debt.get_installment_schedule()
        self.assertTrue(s1[0]['is_paid'])
        self.assertEqual(s1[0]['paid_amount'], '3333333.00')
        self.assertFalse(s1[1]['is_paid'])
        self.assertEqual(s1[1]['paid_amount'], '0.00')

        # Trả kỳ 2: 3.333.333đ
        DebtPayment.create_from_json(debt.pk, self.user, {'amount': '3333333.00'})
        s2 = debt.get_installment_schedule()
        self.assertTrue(s2[0]['is_paid'])
        self.assertTrue(s2[1]['is_paid'])
        self.assertFalse(s2[2]['is_paid'])

        # Trả kỳ 3: 3.333.334đ
        DebtPayment.create_from_json(debt.pk, self.user, {'amount': '3333334.00'})
        s3 = debt.get_installment_schedule()
        self.assertTrue(s3[0]['is_paid'])
        self.assertTrue(s3[1]['is_paid'])
        self.assertTrue(s3[2]['is_paid'])
        self.assertEqual(debt.get_computed_status(), 'FULLY_PAID')

    def test_22_installment_half_payment_allocates_correctly(self):
        """Trả lẻ một nửa kỳ: kỳ đó chưa is_paid, paid_amount đúng."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Mua TiVi',
            'type': 'INSTALLMENT',
            'principal_amount': '6000000.00',
            'installment_months': 3,
            'start_date': '2026-01-01',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # Mỗi kỳ 2.000.000đ. Trả 1.000.000đ (nửa kỳ 1)
        DebtPayment.create_from_json(debt.pk, self.user, {'amount': '1000000.00'})
        s = debt.get_installment_schedule()
        self.assertFalse(s[0]['is_paid'])
        self.assertEqual(s[0]['paid_amount'], '1000000.00')
        self.assertEqual(s[1]['paid_amount'], '0.00')

        # Trả thêm 2.500.000đ (tổng 3.500.000đ -> kỳ 1 đầy 2tr, kỳ 2 có 1.5tr)
        DebtPayment.create_from_json(debt.pk, self.user, {'amount': '2500000.00'})
        s2 = debt.get_installment_schedule()
        self.assertTrue(s2[0]['is_paid'])
        self.assertEqual(s2[0]['paid_amount'], '2000000.00')
        self.assertFalse(s2[1]['is_paid'])
        self.assertEqual(s2[1]['paid_amount'], '1500000.00')
        self.assertFalse(s2[2]['is_paid'])
        self.assertEqual(s2[2]['paid_amount'], '0.00')

    def test_23_installment_single_month(self):
        """installment_months = 1 hoạt động đúng (một kỳ bằng toàn bộ gốc)."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Mua Tu Lanh',
            'type': 'INSTALLMENT',
            'principal_amount': '7500000.00',
            'installment_months': 1,
            'start_date': '2026-01-01',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        schedule = debt.get_installment_schedule()
        self.assertEqual(len(schedule), 1)
        self.assertEqual(schedule[0]['amount'], '7500000.00')
        self.assertEqual(schedule[0]['paid_amount'], '0.00')
        self.assertFalse(schedule[0]['is_paid'])

        # Trả đủ
        DebtPayment.create_from_json(debt.pk, self.user, {'amount': '7500000.00'})
        s_paid = debt.get_installment_schedule()
        self.assertTrue(s_paid[0]['is_paid'])
        self.assertEqual(s_paid[0]['paid_amount'], '7500000.00')
        self.assertEqual(debt.get_computed_status(), 'FULLY_PAID')

    def test_24_lock_order_and_concurrent_payments(self):
        """Khóa Wallet trước DebtLoan và bảo đảm 2 thanh toán đồng thời không vượt dư nợ."""
        import threading
        from django.db import connection

        debt = DebtLoan.create_from_json({
            'person_name': 'Dong Thoi Test',
            'type': 'LOAN_GIVEN',
            'principal_amount': '1000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        results = []
        errors = []

        def pay_task():
            try:
                # Close thread connection to ensure proper isolation
                connection.close()
                p = DebtPayment.create_from_json(debt.pk, self.user, {
                    'amount': '800000.00',
                    'update_wallet': False,
                })
                results.append(p)
            except Exception as e:
                errors.append(e)
            finally:
                connection.close()

        t1 = threading.Thread(target=pay_task)
        t2 = threading.Thread(target=pay_task)
        t1.start()
        t2.start()
        t1.join()
        t2.join()

        debt.refresh_from_db()
        # Tổng thanh toán không bao giờ vượt 1.000.000đ (chỉ 1 cái thành công hoặc SQLite lock serialized)
        self.assertLessEqual(debt.get_total_paid(), Decimal('1000000.00'))

    def test_25_update_payment_exceeding_principal_rejected(self):
        """Sửa payment làm tổng vượt gốc: bị từ chối (400), dữ liệu và ví không đổi."""
        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        debt = DebtLoan.create_from_json({
            'person_name': 'Nguoi Tra',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # Lần 1 trả 3tr, lần 2 trả 1.5tr
        p1 = DebtPayment.create_from_json(debt.pk, self.user, {'amount': '3000000.00', 'update_wallet': True})
        p2 = DebtPayment.create_from_json(debt.pk, self.user, {'amount': '1500000.00', 'update_wallet': True})

        self.wallet.refresh_from_db()
        bal_after_two_payments = self.wallet.get_current_amount()

        # Cố sửa p2 từ 1.5tr lên 2.5tr (tổng 3tr + 2.5tr = 5.5tr > 5tr)
        res = self.client.put(f'/api/debts/{debt.pk}/payments/{p2.pk}', {
            'amount': '2500000.00',
            'update_wallet': True,
        }, format='json')
        self.assertEqual(res.status_code, 400)

        p2.refresh_from_db()
        self.assertEqual(p2.amount, Decimal('1500000.00'))
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), bal_after_two_payments)

    def test_26_payment_without_wallet_update_still_blocked_when_exceeding_remaining(self):
        """Thanh toán update_wallet=False vẫn bị chặn khi vượt dư nợ."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Nguoi Tra Khong Vi',
            'type': 'DEBT_BORROWED',
            'principal_amount': '2000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        # Cố thanh toán 2.500.000đ khi update_wallet=False
        res = self.client.post(f'/api/debts/{debt.pk}/payments/', {
            'amount': '2500000.00',
            'update_wallet': False,
        }, format='json')
        self.assertEqual(res.status_code, 400)
        self.assertEqual(debt.payments.count(), 0)

    def test_27_atomic_rollback_when_exception_occurs_after_wallet_updated(self):
        """Test rollback thật sự: mock exception SAU KHI ví đã update, kiểm tra cả ví và payment đều hoàn tác."""
        debt = DebtLoan.create_from_json({
            'person_name': 'Rollback Partner',
            'type': 'LOAN_GIVEN',
            'principal_amount': '5000000.00',
            'update_wallet': False,
        }, user_pk=self.user.pk)

        self.wallet.refresh_from_db()
        bal_before = self.wallet.get_current_amount()

        # Mock DebtPayment.objects.create để throw exception sau khi ví đã được gọi update_balance
        with mock.patch.object(DebtPayment.objects, 'create', side_effect=RuntimeError('Simulated crash after wallet updated')):
            with self.assertRaises(RuntimeError):
                DebtPayment.create_from_json(debt.pk, self.user, {
                    'amount': '2000000.00',
                    'update_wallet': True,
                })

        self.wallet.refresh_from_db()
        # Ví không được tăng 2.000.000đ mà phải quay về đúng bal_before
        self.assertEqual(self.wallet.get_current_amount(), bal_before)
        # Không có payment nào được lưu
        self.assertEqual(debt.payments.count(), 0)

    def test_28_negative_wallet_balance_behavior_consistent_with_transactions(self):
        """Hệ thống cho phép ví âm: F02 bảo đảm tính toán chính xác với số dư âm giống giao dịch thường."""
        self.wallet.set_balance(Decimal('1000000.00'))  # Ví có 1.000.000đ

        # 1. Cho vay 3.000.000đ (update_wallet=True) -> ví âm -2.000.000đ
        debt = DebtLoan.create_from_json({
            'person_name': 'Ban Vay Am Vi',
            'type': 'LOAN_GIVEN',
            'principal_amount': '3000000.00',
            'update_wallet': True,
        }, user_pk=self.user.pk)

        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), Decimal('-2000000.00'))

        # 2. Thu hồi 500.000đ -> ví tăng lên -1.500.000đ
        DebtPayment.create_from_json(debt.pk, self.user, {
            'amount': '500000.00',
            'update_wallet': True,
        })
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), Decimal('-1500000.00'))

        # 3. Xóa hợp đồng: hoàn lại 3tr ban đầu, trừ lại 500k đã thu -> ví trở về đúng 1.000.000đ ban đầu
        debt.delete()
        self.wallet.refresh_from_db()
        self.assertEqual(self.wallet.get_current_amount(), Decimal('1000000.00'))














